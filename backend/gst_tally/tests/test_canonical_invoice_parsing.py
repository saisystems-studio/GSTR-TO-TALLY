"""Regression tests for the canonical parsing/normalization layer
(services/canonical_invoice.py) -- the single Excel/CSV/JSON -> canonical
schema pipeline shared by GSTR-1, GSTR-2A and GSTR-2B.

Covers:
  * the exact multi-row-header GSTR-2B example from the bug report
  * all 9 (return type x format) combinations producing the same schema
  * the tax-rate rule (explicit column, reliable derivation, else null --
    never a fake 100%)
  * return-type-aware party mapping (never supplier GSTIN under
    customer_gstin's *label*, even though customer_gstin still carries the
    value the Tally voucher pipeline expects)
  * SOURCE_COLUMN_MAPPING_FAILED validation
"""
import json
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import Workbook

from gst_tally.models import GSTInvoice
from gst_tally.services import canonical_invoice as ci


def stream(name, content):
    file_obj = BytesIO(content)
    file_obj.name = name
    return file_obj


class SourceCompanyGstinPriorityTests(TestCase):
    def _csv(self, name, header, rows):
        text = "\n".join([",".join(header), *[",".join(row) for row in rows]])
        return stream(name, text.encode())

    def test_filename_gstin_has_priority_over_file_values(self):
        file_obj = self._csv(
            "APR 2025 returns_12082026_R1_33AFHPM6103Q1Z8_offline_others_0.csv",
            ["Company GSTIN", "Customer GSTIN", "Invoice No", "Invoice Date", "Invoice Value"],
            [["29ABCDE1234F1Z5", "33AAACY4945P1ZS", "INV-1", "01/04/2025", "100"]],
        )
        _, metadata = ci.parse_source(file_obj, "GSTR1", extension="csv")
        self.assertEqual(metadata["company_gstin"], "33AFHPM6103Q1Z8")
        self.assertEqual(metadata["source_company_gstin_source"], "FILENAME")

    def test_dedicated_company_column_is_used_when_filename_has_no_gstin(self):
        file_obj = self._csv(
            "APR_2025_GSTR1.csv",
            ["GSTIN", "Customer GSTIN", "Invoice No", "Invoice Date", "Invoice Value"],
            [["33AFHPM6103Q1Z8", "33AAACY4945P1ZS", "INV-1", "01/04/2025", "100"],
             ["33AFHPM6103Q1Z8", "33AALFC8417G1ZH", "INV-2", "02/04/2025", "200"]],
        )
        _, metadata = ci.parse_source(file_obj, "GSTR1", extension="csv")
        self.assertEqual(metadata["company_gstin"], "33AFHPM6103Q1Z8")
        self.assertEqual(metadata["source_company_gstin_source"], "FILE_GSTIN_COLUMN")

    def test_customer_gstin_is_never_promoted_to_company_gstin(self):
        file_obj = self._csv(
            "APR_2025_GSTR1.csv",
            ["Customer GSTIN", "Invoice No", "Invoice Date", "Invoice Value"],
            [["33AAACY4945P1ZS", "INV-1", "01/04/2025", "100"],
             ["33AALFC8417G1ZH", "INV-2", "02/04/2025", "200"]],
        )
        _, metadata = ci.parse_source(file_obj, "GSTR1", extension="csv")
        self.assertEqual(metadata["company_gstin"], "")
        self.assertEqual(metadata["source_company_gstin_source"], "NOT_FOUND")


class Gstr2bMultiRowHeaderRegressionTests(TestCase):
    """Task spec section 12: the exact multi-row GSTR-2B header/data example."""

    def _workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "GSTR-2B"
        sheet.append(["GSTIN of supplier", "Trade/Legal name", "Invoice Details", None, None, None,
                      "Place of supply", "Supply Attract Reverse Charge", "Taxable Value",
                      "Integrated Tax", "Central Tax", "State/UT Tax", "Cess"])
        sheet.append([None, None, "Invoice number", "Invoice type", "Invoice Date", "Invoice Value",
                      None, None, None, None, None, None, None])
        sheet.append(["33AJJPD4912E1ZQ", "VASANTHAM AGENCIES", "2526/211", "Regular", "07/04/2025", 1680.00,
                      "Tamil Nadu", "No", 1585.96, 0.00, None, None, None])
        buf = BytesIO()
        workbook.save(buf)
        buf.seek(0)
        buf.name = "gstr2b_multirow.xlsx"
        return buf

    def test_two_level_header_is_detected_and_flattened(self):
        rows, metadata = ci.parse_excel(self._workbook(), "GSTR2B")

        diagnostics = metadata["diagnostics"]
        self.assertEqual(diagnostics["detected_header_row"], 1)
        self.assertEqual(diagnostics["detected_subheader_row"], 2)
        self.assertEqual(diagnostics["source_rows"], 1)
        self.assertEqual(diagnostics["normalized_invoice_rows"], 1)
        self.assertEqual(len(rows), 1)

    def test_row_matches_expected_normalized_result(self):
        rows, _ = ci.parse_excel(self._workbook(), "GSTR2B")
        row = rows[0]

        self.assertEqual(row["supplier_gstin"], "33AJJPD4912E1ZQ")
        self.assertEqual(row["supplier_name"], "VASANTHAM AGENCIES")
        self.assertEqual(row["invoice_no"], "2526/211")
        self.assertEqual(row["invoice_type"], "Regular")
        self.assertEqual(row["invoice_date"], date(2025, 4, 7))
        self.assertEqual(row["invoice_value"], Decimal("1680.00"))
        self.assertEqual(row["place_of_supply"], "Tamil Nadu")
        self.assertEqual(row["reverse_charge"], "No")
        self.assertEqual(row["taxable_value"], Decimal("1585.96"))
        self.assertEqual(row["igst"], Decimal("0.00"))

    def test_tax_rate_is_never_the_old_100_percent_bug(self):
        rows, _ = ci.parse_excel(self._workbook(), "GSTR2B")
        self.assertNotEqual(rows[0]["tax_percent"], Decimal("100"))

    def test_customer_gstin_mirrors_supplier_for_tally_voucher_pipeline_compat(self):
        # tally/mappings.py reads customer_gstin as "the counterparty GSTIN"
        # for every return type -- must stay populated for GSTR-2B rows too.
        rows, _ = ci.parse_excel(self._workbook(), "GSTR2B")
        self.assertEqual(rows[0]["customer_gstin"], "33AJJPD4912E1ZQ")

    def test_row_can_construct_a_gstinvoice_without_error(self):
        rows, _ = ci.parse_excel(self._workbook(), "GSTR2B")
        invoice = GSTInvoice(import_batch_id=1, **rows[0])
        self.assertEqual(invoice.supplier_gstin, "33AJJPD4912E1ZQ")


class TaxRateRuleTests(TestCase):
    def test_explicit_rate_column_is_used_directly(self):
        row = ci.normalize_row({"GSTIN of supplier": "33AAAAA0000A1Z5", "Taxable Value": "1000",
                                 "Rate (%)": "18", "Central Tax": "90", "State/UT Tax": "90"}, "GSTR2A")
        self.assertEqual(row["tax_percent"], Decimal("18.000"))

    def test_rate_is_reliably_derived_when_no_explicit_column(self):
        row = ci.normalize_row({"GSTIN of supplier": "33AAAAA0000A1Z5", "Taxable Value": "100",
                                 "Central Tax": "9", "State/UT Tax": "9"}, "GSTR2A")
        self.assertEqual(row["tax_percent"], Decimal("18.000"))

    def test_rate_is_null_when_derivation_is_unreliable(self):
        row = ci.normalize_row({"GSTIN of supplier": "33AAAAA0000A1Z5", "Taxable Value": "1000",
                                 "Central Tax": "47.5", "State/UT Tax": "47.5"}, "GSTR2A")
        self.assertIsNone(row["tax_percent"])

    def test_rate_is_null_never_100_when_no_tax_fields_present_at_all(self):
        row = ci.normalize_row({"GSTIN of supplier": "33AAAAA0000A1Z5", "Taxable Value": "1000"}, "GSTR2A")
        self.assertIsNone(row["tax_percent"])

    def test_genuinely_zero_tax_derives_zero_not_none_or_100(self):
        row = ci.normalize_row({"GSTIN of supplier": "33AAAAA0000A1Z5", "Taxable Value": "1000",
                                 "Integrated Tax": "0"}, "GSTR2A")
        self.assertEqual(row["tax_percent"], Decimal("0.000"))


class ReturnTypePartyMappingTests(TestCase):
    def test_gstr1_uses_customer_fields(self):
        row = ci.normalize_row({"GSTIN/UIN of Recipient": "29BBBBB1111B1Z5", "Receiver Name": "Acme Buyer",
                                 "Invoice Number": "IN-1", "Taxable Value": "500"}, "GSTR1")
        self.assertEqual(row["customer_gstin"], "29BBBBB1111B1Z5")
        self.assertEqual(row["customer_name"], "Acme Buyer")
        self.assertEqual(row["supplier_gstin"], "")
        self.assertEqual(row["supplier_name"], "")

    def test_gstr2a_uses_supplier_fields(self):
        row = ci.normalize_row({"Supplier GSTIN": "27CCCCC2222C1Z5", "Supplier Name": "Vendor Co",
                                 "Invoice No": "IN-2", "Taxable Value": "500"}, "GSTR2A")
        self.assertEqual(row["supplier_gstin"], "27CCCCC2222C1Z5")
        self.assertEqual(row["supplier_name"], "Vendor Co")
        self.assertEqual(row["customer_name"], "")

    def test_gstr2b_uses_supplier_fields(self):
        row = ci.normalize_row({"GSTIN of supplier": "27CCCCC2222C1Z5", "Trade/Legal name": "Vendor Co",
                                 "Invoice number": "IN-2", "Taxable Value": "500"}, "GSTR2B")
        self.assertEqual(row["supplier_gstin"], "27CCCCC2222C1Z5")
        self.assertEqual(row["supplier_name"], "Vendor Co")


class NineCombinationTests(TestCase):
    """All 9 (return type x format) combinations produce the same canonical
    schema, per the task's acceptance criteria."""

    REQUIRED_KEYS = {"invoice_date", "supplier_gstin", "supplier_name", "customer_gstin", "customer_name",
                      "invoice_no", "invoice_type", "invoice_value", "taxable_value", "tax_percent",
                      "cgst", "sgst", "igst", "cess", "place_of_supply", "state_code", "reverse_charge"}

    def _assert_canonical_shape(self, row):
        self.assertTrue(self.REQUIRED_KEYS.issubset(row.keys()))

    # ---- GSTR-1 ----
    def test_gstr1_excel(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Recipient GSTIN", "Customer Name", "Invoice No", "Invoice Date", "Taxable Value",
                      "Rate", "Integrated Tax", "Invoice Value"])
        sheet.append(["29AAAAA0000A1Z5", "Acme Buyer", "INV-1", "10/05/2025", 1000, 18, 180, 1180])
        buf = BytesIO()
        workbook.save(buf)
        buf.seek(0)
        buf.name = "gstr1.xlsx"
        rows, _ = ci.parse_excel(buf, "GSTR1")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["customer_gstin"], "29AAAAA0000A1Z5")
        self.assertEqual(rows[0]["customer_name"], "Acme Buyer")
        self.assertEqual(rows[0]["tax_percent"], Decimal("18.000"))

    def test_gstr1_csv(self):
        text = ("Recipient GSTIN,Customer Name,Invoice No,Invoice Date,Taxable Value,Rate,Integrated Tax,Invoice Value\n"
                "29AAAAA0000A1Z5,Acme Buyer,INV-1,10/05/2025,1000,18,180,1180\n")
        rows, _ = ci.parse_csv(stream("gstr1.csv", text.encode()), "GSTR1")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["customer_gstin"], "29AAAAA0000A1Z5")

    def test_gstr1_json(self):
        payload = {"gstin": "33ZZZZZ9999Z1Z5", "fp": "042025", "b2b": [{"ctin": "29AAAAA0000A1Z5", "inv": [{
            "inum": "INV-1", "idt": "10-05-2025", "val": 1180, "pos": "29",
            "itms": [{"itm_det": {"txval": 1000, "rt": 18, "camt": 0, "samt": 0, "iamt": 180}}]}]}]}
        rows, _ = ci.parse_json(stream("gstr1.json", json.dumps(payload).encode()), "GSTR1")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["customer_gstin"], "29AAAAA0000A1Z5")

    # ---- GSTR-2A ----
    def test_gstr2a_excel(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["GSTIN of supplier", "Trade/Legal name", "Invoice number", "Invoice Date",
                      "Taxable Value", "Rate", "Central Tax", "State/UT Tax", "Invoice Value"])
        sheet.append(["33AJJPD4912E1ZQ", "VASANTHAM AGENCIES", "2526/211", "07/04/2025", 1585.96, 18, 142.74, 142.74, 1871.44])
        buf = BytesIO()
        workbook.save(buf)
        buf.seek(0)
        buf.name = "gstr2a.xlsx"
        rows, _ = ci.parse_excel(buf, "GSTR2A")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["supplier_gstin"], "33AJJPD4912E1ZQ")
        self.assertEqual(rows[0]["supplier_name"], "VASANTHAM AGENCIES")

    def test_gstr2a_csv(self):
        text = ("GSTIN of supplier,Trade/Legal name,Invoice number,Invoice Date,Taxable Value,Rate,Central Tax,State/UT Tax,Invoice Value\n"
                "33AJJPD4912E1ZQ,VASANTHAM AGENCIES,2526/211,07/04/2025,1585.96,18,142.74,142.74,1871.44\n")
        rows, _ = ci.parse_csv(stream("gstr2a.csv", text.encode()), "GSTR2A")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["supplier_gstin"], "33AJJPD4912E1ZQ")

    def test_gstr2a_json(self):
        payload = {"b2b": [{"ctin": "33AJJPD4912E1ZQ", "inv": [{
            "inum": "2526/211", "idt": "07-04-2025", "val": 1871.44, "pos": "33",
            "itms": [{"itm_det": {"txval": 1585.96, "rt": 18, "camt": 142.74, "samt": 142.74}}]}]}]}
        rows, _ = ci.parse_json(stream("gstr2a.json", json.dumps(payload).encode()), "GSTR2A")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["supplier_gstin"], "33AJJPD4912E1ZQ")

    # ---- GSTR-2B ----
    def test_gstr2b_excel(self):
        rows, _ = ci.parse_excel(Gstr2bMultiRowHeaderRegressionTests()._workbook(), "GSTR2B")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["supplier_gstin"], "33AJJPD4912E1ZQ")

    def test_gstr2b_csv(self):
        text = ("GSTIN of supplier,Trade/Legal name,Invoice number,Invoice type,Invoice Date,Taxable Value,Integrated Tax,Invoice Value,Place of supply\n"
                "33AJJPD4912E1ZQ,VASANTHAM AGENCIES,2526/211,Regular,07/04/2025,1585.96,0,1585.96,Tamil Nadu\n")
        rows, _ = ci.parse_csv(stream("gstr2b.csv", text.encode()), "GSTR2B")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["supplier_gstin"], "33AJJPD4912E1ZQ")
        self.assertEqual(rows[0]["place_of_supply"], "Tamil Nadu")

    def test_gstr2b_json(self):
        payload = {"b2b": [{"ctin": "33AJJPD4912E1ZQ", "inv": [{
            "inum": "2526/211", "idt": "07-04-2025", "val": 1585.96, "pos": "33",
            "itms": [{"itm_det": {"txval": 1585.96, "rt": 0, "iamt": 0}}]}]}]}
        rows, _ = ci.parse_json(stream("gstr2b.json", json.dumps(payload).encode()), "GSTR2B")
        self.assertEqual(len(rows), 1)
        self._assert_canonical_shape(rows[0])
        self.assertEqual(rows[0]["supplier_gstin"], "33AJJPD4912E1ZQ")


class FormatReturnTypeIndependenceTests(TestCase):
    """The parser dispatch bug: format (how a file is read) must never
    determine which return type's alias/party mapping is applied -- a GSTR-1
    file uploaded as Excel must still use customer-side aliases, not
    GSTR-2B's supplier-side ones."""

    def test_gstr1_uploaded_as_excel_uses_customer_aliases(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Recipient GSTIN", "Customer Name", "Invoice No", "Invoice Date", "Taxable Value", "Invoice Value"])
        sheet.append(["29AAAAA0000A1Z5", "Acme Buyer", "INV-9", "01/01/2025", 1000, 1180])
        buf = BytesIO()
        workbook.save(buf)
        buf.seek(0)
        buf.name = "gstr1.xlsx"
        rows, _ = ci.parse_source(buf, "GSTR1", extension="xlsx")
        self.assertEqual(rows[0]["customer_gstin"], "29AAAAA0000A1Z5")
        self.assertEqual(rows[0]["supplier_gstin"], "")


class SourceColumnMappingFailedTests(TestCase):
    def test_unmapped_columns_raise_a_clear_error_instead_of_silently_importing_blanks(self):
        text = "Invoice Number,Filing Period\n" + "\n".join(f",FY2025-{i:02d}" for i in range(1, 6)) + "\n"
        with self.assertRaisesRegex(ValueError, "SOURCE_COLUMN_MAPPING_FAILED"):
            ci.parse_csv(stream("bad.csv", text.encode()), "GSTR2A")

    def test_a_small_file_with_no_matches_does_not_false_positive(self):
        # Below the row-count threshold -- must not raise for tiny/edge files,
        # even though this one row's Invoice Number cell is blank.
        text = "Invoice Number,Filing Period\n,FY2025-01\n"
        rows, _ = ci.parse_csv(stream("small.csv", text.encode()), "GSTR2A")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["invoice_no"], "")
