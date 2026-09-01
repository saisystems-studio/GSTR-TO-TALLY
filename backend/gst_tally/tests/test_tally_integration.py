from decimal import Decimal
from datetime import date
from io import BytesIO
from unittest.mock import patch
import json
from xml.etree import ElementTree as ET
from django.test import SimpleTestCase
from django.test import override_settings
from gst_tally.tally.connection import connection_status, step3_connection_check
from gst_tally.tally.odbc import odbc_company_status, verify_source_company
from gst_tally.tally.client import TallyClient, TallyConnectionError
from gst_tally.tally.mappings import sales_ledger_name
from gst_tally.tally.response_parser import TallyResponse, parse_response
from gst_tally.tally.validators import transaction_type, validate_voucher
from gst_tally.tally.voucher_builder import build_voucher
from gst_tally.tally.json_voucher_builder import build_json_voucher
from gst_tally.tally.json_master_builder import build_json_master
from gst_tally.tally.json_response_parser import parse_json_response
from gst_tally.tally.json_company_reader import get_tally_company_gst_registration
from gst_tally.tally.read_parsers import parse_voucher_query_response
from gst_tally.tally.voucher_verification import build_day_book_request, verify_voucher
from gst_tally.services.company import financial_year_details, get_financial_year_start, normalize_company, source_company_metadata
from gst_tally.utils.file_utils import parse_date
from gst_tally.services.party_lookup import party_eligibility
from gst_tally.services.gstr1_parser import parse as parse_gstr1
from gst_tally.services.gstr2b_parser import parse as parse_gstr2b
from gst_tally.services.gstr2a_parser import parse as parse_gstr2a
from gst_tally.tally.service import (_period_payload, _voucher_balance, _voucher_diagnostics,
                                     _write_error_code, _query_back_state, _verified_write_outcome,
                                     _write_voucher, _write_master, _write_metadata, _query_preflight_error)
from openpyxl import Workbook


def voucher(rate, transaction="INTRA-STATE"):
    tax = Decimal("1000") * Decimal(str(rate)) / 100
    cgst = tax / 2 if transaction == "INTRA-STATE" else Decimal("0")
    sgst = cgst
    igst = tax if transaction == "INTER-STATE" else Decimal("0")
    return {"invoice_number": "INV-1", "invoice_date": "2026-08-22", "party": {"name": "Buyer", "gstin": "33AAACB2894G1ZJ"},
            "items": [{"taxable_value": "1000", "gst_rate": str(rate)}], "transaction_type": transaction,
            "taxable_total": "1000", "cgst": str(cgst), "sgst": str(sgst), "igst": str(igst), "cess": "0",
            "other_charges": "0", "round_off": "0", "invoice_total": str(Decimal("1000") + tax)}


class TallyValidationTests(SimpleTestCase):
    def test_voucher_parser_extracts_nested_purchase_and_tax_ledger_gst_details(self):
        raw = b'''<ENVELOPE><BODY><DATA><COLLECTION><VOUCHER MASTERID="411">
          <DATE>20250401</DATE><VOUCHERNUMBER>1</VOUCHERNUMBER><VOUCHERTYPENAME>Purchase</VOUCHERTYPENAME>
          <REFERENCE>1</REFERENCE><PARTYLEDGERNAME>Supplier</PARTYLEDGERNAME>
          <LEDGERENTRIES.LIST><LEDGERNAME>GST Purchase 18%</LEDGERNAME><AMOUNT>-1000</AMOUNT>
            <GSTDETAILS.LIST><TAXABILITY>Taxable</TAXABILITY><RATEDETAILS.LIST><GSTRATEDUTYHEAD>Central Tax</GSTRATEDUTYHEAD><GSTRATE>9</GSTRATE></RATEDETAILS.LIST><RATEDETAILS.LIST><GSTRATEDUTYHEAD>State Tax</GSTRATEDUTYHEAD><GSTRATE>9</GSTRATE></RATEDETAILS.LIST></GSTDETAILS.LIST>
          </LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>GST Purchase 5%</LEDGERNAME><AMOUNT>-500</AMOUNT>
            <GSTDETAILS.LIST><TAXABILITY>Taxable</TAXABILITY><RATEDETAILS.LIST><GSTRATEDUTYHEAD>Central Tax</GSTRATEDUTYHEAD><GSTRATE>2.5</GSTRATE></RATEDETAILS.LIST><RATEDETAILS.LIST><GSTRATEDUTYHEAD>State Tax</GSTRATEDUTYHEAD><GSTRATE>2.5</GSTRATE></RATEDETAILS.LIST></GSTDETAILS.LIST>
          </LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>Input CGST 9%</LEDGERNAME><AMOUNT>-90</AMOUNT><GSTDUTYHEAD>Central Tax</GSTDUTYHEAD><RATEOFINVOICETAX.LIST><RATEOFINVOICETAX>9</RATEOFINVOICETAX></RATEOFINVOICETAX.LIST></LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>Input SGST 9%</LEDGERNAME><AMOUNT>-90</AMOUNT><GSTDUTYHEAD>State Tax</GSTDUTYHEAD><RATEOFINVOICETAX.LIST><RATEOFINVOICETAX>9</RATEOFINVOICETAX></RATEOFINVOICETAX.LIST></LEDGERENTRIES.LIST>
        </VOUCHER></COLLECTION></DATA></BODY></ENVELOPE>'''
        result = parse_voucher_query_response(raw)
        entries = {row["ledger"]: row for row in result["vouchers"][0]["ledger_entries"]}
        self.assertEqual(entries["GST Purchase 18%"]["gst_rate"], "18")
        self.assertEqual(entries["GST Purchase 5%"]["gst_rate"], "5")
        self.assertEqual(entries["GST Purchase 18%"]["amount"], "-1000")
        self.assertEqual(entries["Input CGST 9%"]["gst_classification"], "Central Tax")
        self.assertIn("GSTDETAILS.LIST", entries["GST Purchase 18%"]["rate_detail_fields"])

    def test_purchase_query_back_recognizes_multiple_rates_and_tax_ledgers(self):
        raw = b'''<ENVELOPE><BODY><DATA><COLLECTION><VOUCHER MASTERID="411"><DATE>20250401</DATE>
          <VOUCHERNUMBER>1</VOUCHERNUMBER><VOUCHERTYPENAME>Purchase</VOUCHERTYPENAME><REFERENCE>1</REFERENCE><PARTYLEDGERNAME>Supplier</PARTYLEDGERNAME>
          <LEDGERENTRIES.LIST><LEDGERNAME>GST Purchase 18%</LEDGERNAME><AMOUNT>-1000</AMOUNT><RATEOFINVOICETAX.LIST><RATEOFINVOICETAX>18</RATEOFINVOICETAX></RATEOFINVOICETAX.LIST></LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>GST Purchase 5%</LEDGERNAME><AMOUNT>-500</AMOUNT><RATEOFINVOICETAX.LIST><RATEOFINVOICETAX>5</RATEOFINVOICETAX></RATEOFINVOICETAX.LIST></LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>Input CGST 9%</LEDGERNAME><AMOUNT>-90</AMOUNT></LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>Input SGST 9%</LEDGERNAME><AMOUNT>-90</AMOUNT></LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>Input CGST 2.5%</LEDGERNAME><AMOUNT>-12.5</AMOUNT></LEDGERENTRIES.LIST>
          <LEDGERENTRIES.LIST><LEDGERNAME>Input SGST 2.5%</LEDGERNAME><AMOUNT>-12.5</AMOUNT></LEDGERENTRIES.LIST>
        </VOUCHER></COLLECTION></DATA></BODY></ENVELOPE>'''
        class Client:
            def post(self, payload): return raw
        voucher = {"invoice_number": "1", "invoice_date": "2025-04-01", "voucher_type": "Purchase",
                   "party": {"name": "Supplier"},
                   "rate_allocations": [{"account_ledger": "GST Purchase 18%", "gst_rate": "18", "taxable_value": "1000"},
                                        {"account_ledger": "GST Purchase 5%", "gst_rate": "5", "taxable_value": "500"}],
                   "cgst": "102.5", "sgst": "102.5", "igst": "0"}
        result = verify_voucher(Client(), "SRI MAHALAKSHMI TRADERS", voucher)
        self.assertTrue(result["found"])
        self.assertEqual(result["actual_gst_rates"], ["5", "18"])
        self.assertEqual(set(result["actual_tax_ledgers"]), {"Input CGST 9%", "Input SGST 9%", "Input CGST 2.5%", "Input SGST 2.5%"})

    def test_query_back_rejects_ledger_master_rate_when_allocation_rate_is_blank(self):
        raw = b'''<ENVELOPE><BODY><DATA><COLLECTION><VOUCHER MASTERID="411"><DATE>20250401</DATE>
          <VOUCHERNUMBER>1</VOUCHERNUMBER><VOUCHERTYPENAME>Purchase</VOUCHERTYPENAME><REFERENCE>1</REFERENCE>
          <LEDGERENTRIES.LIST><LEDGERNAME>Domestic Purchases A</LEDGERNAME><AMOUNT>-1000</AMOUNT></LEDGERENTRIES.LIST>
        </VOUCHER></COLLECTION></DATA></BODY></ENVELOPE>'''
        class Client(TallyClient):
            def post(self, payload, headers=None): return raw
        voucher = {"invoice_number": "1", "invoice_date": "2025-04-01", "voucher_type": "Purchase", "party": {},
                   "rate_allocations": [{"account_ledger": "Domestic Purchases A", "gst_rate": "18", "taxable_value": "1000"}],
                   "tax_allocations": []}
        with patch("gst_tally.tally.connection.ledger_details",
                   return_value={"exists": True, "gst_rate": "18", "taxability": "Taxable"}):
            result = verify_voucher(Client(), "SRI MAHALAKSHMI TRADERS", voucher)
        self.assertFalse(result["found"])
        entry = result["raw_tally_ledger_fields"][0]
        self.assertEqual(entry["gst_rate"], "")
        self.assertEqual(entry["gst_rate_source"], "")

    def test_excel_customer_gstin_precedes_generic_company_gstin_and_deduplicates_later(self):
        gstins = ["33ABFFA3666C1ZU", "33AGHPV2166P1ZS", "33AAACY4945P1ZS",
                  "33FRWPM7826N1ZP", "33BVBPK4067N2ZC", "33AADCR3002K1ZS"]
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["GSTIN", "Customer GSTIN", "Invoice number", "Invoice Date", "Taxable Value", "Rate", "CGST", "SGST", "Invoice Value"])
        for index, gstin in enumerate(gstins, 1):
            sheet.append(["33AFHPM6103Q1Z8", gstin, str(index), "01-04-2025", 100, 18, 9, 9, 118])
            if index in {1, 2, 4, 5}: sheet.append(["33AFHPM6103Q1Z8", gstin, str(index), "01-04-2025", 50, 28, 7, 7, 64])
        stream = BytesIO(); workbook.save(stream); stream.seek(0)
        rows, _ = parse_gstr2b(stream)
        self.assertEqual({row["customer_gstin"] for row in rows}, set(gstins))
        self.assertEqual(sum(row["customer_gstin"] == gstins[0] for row in rows), 2)
        self.assertNotIn("33AFHPM6103Q1Z8", {row["customer_gstin"] for row in rows})

    def test_day_book_verification_uses_no_custom_tdl_or_collection(self):
        xml = build_day_book_request("GSTRCOMPANY", "2026-07-24").decode()
        for forbidden in ("<TDL>", "<TDLMESSAGE>", "<COLLECTION", "<TYPE>Collection", "<FETCH>"):
            self.assertNotIn(forbidden, xml)
        self.assertIn("<REPORTNAME>Day Book</REPORTNAME>", xml)

    def test_standard_day_book_response_confirms_sales_voucher(self):
        class Client:
            def post(self, payload):
                if b"<REPORTNAME>Day Book</REPORTNAME>" in payload:
                    return b'<ENVELOPE><BODY><DATA><DSPVCHDETAILS><DSPVCHNUMBER>2627B2B118</DSPVCHNUMBER><DSPVCHTYPE>Sales</DSPVCHTYPE><DSPVCHLEDACCOUNT>N.K.R. MALIGAI</DSPVCHLEDACCOUNT><VOUCHERID>4</VOUCHERID></DSPVCHDETAILS></DATA></BODY></ENVELOPE>'
                return b'<ENVELOPE><HEADER><STATUS>1</STATUS></HEADER><BODY><DESC></DESC><DATA><COLLECTION></COLLECTION></DATA></BODY></ENVELOPE>'
        result = verify_voucher(Client(), "GSTRCOMPANY", {"invoice_number": "2627B2B118", "invoice_date": "2026-07-24", "party": {"name": "N.K.R. MALIGAI"}})
        self.assertTrue(result["found"])
        self.assertEqual(result["identifier"], "4")
        self.assertIn("DAY_BOOK", result["source"])

    def test_envelope_without_export_payload_is_not_an_authoritative_voucher_query(self):
        class Client:
            def post(self, payload):
                return b'<ENVELOPE><HEADER><TALLYREQUEST>Import Data</TALLYREQUEST></HEADER><BODY><IMPORTDATA><REQUESTDESC><REPORTNAME>All Masters</REPORTNAME></REQUESTDESC></IMPORTDATA></BODY></ENVELOPE>'

        result = verify_voucher(Client(), "SRI MAHALAKSHMI TRADERS,", {
            "invoice_number": "1", "invoice_date": "2025-04-01",
            "voucher_type": "Purchase", "party": {"name": "Supplier"},
        })

        self.assertFalse(result["query_valid"])
        self.assertFalse(result["found"])
        self.assertIn("carries no export payload", result["reason"])

    def test_query_back_state_distinguishes_found_missing_and_invalid(self):
        self.assertEqual(_query_back_state({"query_valid": True, "found": True}), "FOUND")
        self.assertEqual(_query_back_state({"query_valid": True, "found": False}), "MISSING")
        self.assertEqual(_query_back_state({"query_valid": False, "found": False}), "INVALID")

    def test_invalid_query_response_blocks_write_preflight(self):
        self.assertEqual(_query_preflight_error({"query_valid": True, "found": False}), "")
        self.assertIn("query preflight failed", _query_preflight_error({
            "query_valid": False, "found": False, "reason": "wrong report",
        }))

    def test_accepted_write_without_positive_query_back_is_not_imported(self):
        response = TallyResponse(created=1, errors=0)

        missing = _verified_write_outcome(response, {"query_valid": True, "found": False, "reason": "not found"})
        invalid = _verified_write_outcome(response, {"query_valid": False, "found": False, "reason": "wrong response"})
        found = _verified_write_outcome(response, {"query_valid": True, "found": True, "reason": "found"})

        self.assertEqual(missing["result_status"], "Verification Failed")
        self.assertEqual(missing["mapping_status"], "Unknown")
        self.assertEqual(invalid["result_status"], "Verification Failed")
        self.assertEqual(invalid["mapping_status"], "Unknown")
        self.assertEqual(found["result_status"], "Imported")
        self.assertEqual(found["mapping_status"], "Imported")

    @override_settings(TALLY_WRITE_FORMAT="JSON")
    def test_active_writes_use_native_json_for_masters_and_xml_for_vouchers(self):
        """Verified against a live TallyPrime instance: master writes (Create/
        Alter of a ledger) must go over Tally's native JSON import
        (import_json), carrying the \\x04-prefixed fixed-list enums Tally's
        own export uses -- never XML. Voucher writes stay on the existing XML
        import path (import_data); this split is unaffected by
        TALLY_WRITE_FORMAT, which is informational metadata only."""
        class Client:
            def import_data(self, payload):
                self.payloads = getattr(self, "payloads", []) + [payload]
                return TallyResponse(created=1)
            def import_json(self, payload, object_id):
                self.json_payloads = getattr(self, "json_payloads", []) + [payload]
                return TallyResponse(created=1)

        value = voucher(18)
        value["items"][0]["sales_ledger"] = "GST Sales 18%"
        client = Client()
        _write_voucher(client, value, "D")
        _write_master(client, {"master_type": "Charge", "name": "Round Off", "group": "Indirect Expenses"}, "D")

        self.assertEqual(len(client.payloads), 1)
        self.assertTrue(client.payloads[0].startswith(b"<ENVELOPE>"))
        self.assertEqual(len(client.json_payloads), 1)
        master_message = client.json_payloads[0]["tallymessage"][0]
        self.assertEqual(master_message["metadata"]["name"], "Round Off")
        metadata = _write_metadata()
        self.assertEqual(metadata["active_write_format"], "JSON_MASTERS_XML_VOUCHERS")
        self.assertTrue(metadata["xml_used_for_active_voucher_write"])
        self.assertTrue(metadata["json_used_for_active_master_write"])
        self.assertFalse(metadata["xml_used_for_active_master_write"])
        self.assertTrue(metadata["write_enabled"])

    def test_source_cgst_sgst_are_preserved_in_voucher_xml(self):
        value = voucher(5); value["cgst"], value["sgst"], value["igst"] = "25.01", "24.99", "0"
        value["invoice_total"] = "1050"; value["items"][0]["sales_ledger"] = "GST SALES @ 5%"
        xml = ET.fromstring(build_voucher(value, "GSTRCOMPANY"))
        entries = {node.findtext("LEDGERNAME"): node.findtext("AMOUNT") for node in xml.findall(".//LEDGERENTRIES.LIST")}
        self.assertEqual(entries["CGST"], "25.01"); self.assertEqual(entries["SGST"], "24.99"); self.assertNotIn("IGST", entries)
        self.assertEqual(xml.findtext(".//VOUCHER/ISINVOICE"), "Yes")
        self.assertEqual(xml.find(".//VOUCHER").get("OBJVIEW"), "Invoice Voucher View")
        self.assertEqual(xml.findtext(".//VOUCHER/PERSISTEDVIEW"), "Invoice Voucher View")
        self.assertIsNone(xml.find(".//ALLINVENTORYENTRIES.LIST"))

    def test_source_igst_is_preserved_without_generated_cgst_sgst(self):
        value = voucher(5, "INTER-STATE"); value["igst"] = "50.00"; value["items"][0]["sales_ledger"] = "GST SALES @ 5%"
        xml = ET.fromstring(build_voucher(value, "GSTRCOMPANY"))
        entries = {node.findtext("LEDGERNAME"): node.findtext("AMOUNT") for node in xml.findall(".//LEDGERENTRIES.LIST")}
        self.assertEqual(entries["IGST"], "50.00"); self.assertNotIn("CGST", entries); self.assertNotIn("SGST", entries)

    def test_only_one_intra_state_component_requires_review(self):
        value = voucher(5); value.update(cgst="50", sgst="0", invoice_total="1050")
        result = validate_voucher(value)
        self.assertTrue(result["review_required"])
        self.assertIn("Source GST breakup is incomplete/inconsistent.", result["review_reasons"])

    def test_mixed_intra_and_inter_state_components_require_review(self):
        value = voucher(5); value.update(igst="1", invoice_total="1051")
        result = validate_voucher(value)
        self.assertTrue(result["review_required"])
        self.assertIn("Both intra-state and inter-state GST components are present.", result["review_reasons"])

    def test_difference_is_diagnostic_and_round_off_is_ignored(self):
        value = voucher(5, "INTER-STATE"); value.update(taxable_total="23828.70", igst="1191.44", round_off="0", invoice_total="25020.00")
        value["items"] = [{"taxable_value": "23828.70", "gst_rate": "5"}]
        result = validate_voucher(value)
        self.assertTrue(result["critical_valid"])
        self.assertEqual(result["calculated_total"], "25020.14")
        self.assertEqual(result["difference"], "0.14")
        self.assertEqual(value["igst"], "1191.44")

    def test_source_total_difference_does_not_manufacture_balancing_ledger(self):
        value = voucher(5, "INTER-STATE")
        value.update(taxable_total="23828.70", igst="1191.44", invoice_total="25020.00")
        value["items"] = [{"taxable_value": "23828.70", "gst_rate": "5", "sales_ledger": "GST SALES @ 5%"}]
        value["validation"] = validate_voucher(value)
        xml = ET.fromstring(build_voucher(value, "GSTRCOMPANY"))
        entries = {node.findtext("LEDGERNAME"): node.findtext("AMOUNT") for node in xml.findall(".//LEDGERENTRIES.LIST")}
        self.assertEqual(entries["IGST"], "1191.44")
        self.assertNotIn("Source Total Adjustment", entries)
        self.assertNotIn("CGST", entries)
        self.assertNotIn("SGST", entries)

    def test_multi_rate_invoice_builds_one_voucher_with_sorted_sales_allocations(self):
        value = voucher(18)
        value.update(invoice_number="32", invoice_total="153800.16", taxable_total="125837.00",
                     cgst="13981.58", sgst="13981.58", igst="0")
        value["items"] = [
            {"taxable_value": "53125.00", "gst_rate": "28", "sales_ledger": "GST SALES @ 28%"},
            {"taxable_value": "72712.00", "gst_rate": "18", "sales_ledger": "GST SALES @ 18%"},
        ]
        value["rate_allocations"] = [
            {"gst_rate": "28", "sales_ledger": "GST SALES @ 28%", "taxable_value": "53125.00", "cgst": "7437.50", "sgst": "7437.50", "igst": "0"},
            {"gst_rate": "18", "sales_ledger": "GST SALES @ 18%", "taxable_value": "72712.00", "cgst": "6544.08", "sgst": "6544.08", "igst": "0"},
        ]
        xml = ET.fromstring(build_voucher(value, "GSTRCOMPANY"))
        self.assertEqual(len(xml.findall(".//VOUCHER")), 1)
        entries = xml.findall(".//LEDGERENTRIES.LIST")
        names = [entry.findtext("LEDGERNAME") for entry in entries]
        self.assertEqual(names, ["Buyer", "GST SALES @ 18%", "GST SALES @ 28%", "CGST", "SGST"])
        self.assertEqual(sum(entry.findtext("ISPARTYLEDGER") == "Yes" for entry in entries), 1)
        amounts = {entry.findtext("LEDGERNAME"): entry.findtext("AMOUNT") for entry in entries}
        self.assertEqual(amounts["GST SALES @ 18%"], "72712.00")
        self.assertEqual(amounts["GST SALES @ 28%"], "53125.00")

    def test_native_json_multi_rate_accounting_voucher_uses_official_envelope(self):
        value = voucher(18); value.update(invoice_number="32", invoice_date="2025-04-17", invoice_total="153800.16",
                                          taxable_total="125837.00", cgst="13981.58", sgst="13981.58")
        value["rate_allocations"] = [
            {"gst_rate": "28", "sales_ledger": "GST SALES @ 28%", "taxable_value": "53125.00"},
            {"gst_rate": "18", "sales_ledger": "GST SALES @ 18%", "taxable_value": "72712.00"},
        ]
        payload = build_json_voucher(value, "Company D")
        self.assertEqual(payload["static_variables"], [{"name": "svVchImportFormat", "value": "jsonex"},
                         {"name": "svCurrentCompany", "value": "Company D"},
                         {"name": "svFromDate", "value": "20250401"},
                         {"name": "svToDate", "value": "20260331"}])
        message = payload["tallymessage"][0]
        self.assertEqual(message["metadata"], {"type": "Voucher", "vchtype": "Sales", "action": "Create", "objview": "Invoice Voucher View"})
        self.assertEqual(message["date"], "20250417"); self.assertTrue(message["isinvoice"])
        self.assertEqual(message["persistedview"], "Invoice Voucher View")
        self.assertNotIn("allinventoryentries", message); self.assertNotIn("roundoff", message)
        self.assertEqual([row["ledgername"] for row in message["ledgerentries"]],
                         ["Buyer", "GST SALES @ 18%", "GST SALES @ 28%", "CGST", "SGST"])

    def test_native_json_master_and_response_follow_tally_jsonex_contract(self):
        payload = build_json_master({"master_type": "Party", "name": "Buyer", "group": "Sundry Debtors",
                                     "gstin": "33AAACB2894G1ZJ", "state": "Tamil Nadu", "pincode": "625001"}, "Company D")
        self.assertEqual(payload["static_variables"][0], {"name": "svMstImportFormat", "value": "jsonex"})
        self.assertEqual(payload["tallymessage"][0]["metadata"]["type"], "Ledger")
        response = parse_json_response(b'{"status":"1","data":{"import_result":{"created":1,"errors":0,"exceptions":0,"lastvchid":70}}}')
        self.assertTrue(response.accepted); self.assertEqual(response.last_vch_id, "70")
        rejected = parse_json_response(b'{"status":"0","data":{"import_result":{"created":0,"errors":1,"description":"Invalid voucher"}}}')
        self.assertFalse(rejected.accepted); self.assertEqual(rejected.error, "Invalid voucher")

    def test_native_json_client_sends_required_import_headers(self):
        class Client(TallyClient):
            def post(self, payload, headers=None): self.sent = (payload, headers); return b'{"status":"1","data":{"import_result":{"created":1}}}'
        client = Client(); response = client.import_json({"tallymessage": []}, "Vouchers")
        self.assertTrue(response.accepted)
        self.assertEqual(client.sent[1], {"Content-Type": "application/json", "version": "1", "tallyrequest": "Import", "type": "Data", "id": "Vouchers", "detailed-response": "Yes"})

    def test_native_json_company_reader_matches_expected_gstin_across_multiple_tax_units(self):
        class Client:
            def post(self, payload, headers=None):
                subtype, object_id = headers.get("subtype"), headers.get("id")
                if subtype == "Company":
                    return json.dumps({"status": "1", "prod_maj_rel": "7", "prod_min_rel": "1", "tallymessage": [{
                        "statename": {"value": "Tamil Nadu"}, "isgston": {"value": True}, "exciseunitname": {"value": "TN Registration"},
                        "booksfrom": {"value": "20250401"}, "startingfrom": {"value": "20250401"},
                        "gstregistrationdetails": [{"gstregnumber": {"value": "33AFHPM6103Q1Z8"}},
                                                   {"gstregnumber": {"value": "32AAACB2894G1ZK"}}]}]}).encode()
                gstin = "33AFHPM6103Q1Z8" if object_id == "TN Registration" else "32AAACB2894G1ZK"
                return json.dumps({"status": "1", "tallymessage": [{"gstregnumber": {"value": gstin},
                    "statename": {"value": "Tamil Nadu"}, "gstregistrationdetails": [{"registrationtype": {"value": "Regular"}}]}]}).encode()
        result = get_tally_company_gst_registration("D", "33AFHPM6103Q1Z8", Client())
        self.assertEqual(result["gstin"], "33AFHPM6103Q1Z8"); self.assertEqual(result["gst_registrations_found"], 2)
        self.assertEqual(result["gstin_read_source"], "JSON_GST_REGISTRATION")
        self.assertEqual(result["tally_version"], "7.1"); self.assertEqual(result["books_from"], date(2025, 4, 1))

    def test_company_reader_never_sends_collection_or_tdl(self):
        class Client:
            def __init__(self): self.requests = []
            def post(self, payload, headers=None):
                self.requests.append((json.loads(payload), headers))
                if headers["subtype"] == "Company":
                    return b'{"status":"1","tallymessage":[{"exciseunitname":{"value":"Default Tax Unit"}}]}'
                return b'{"status":"1","tallymessage":[{}]}'
        client = Client(); get_tally_company_gst_registration("D", "33AFHPM6103Q1Z8", client)
        self.assertTrue(client.requests)
        for payload, headers in client.requests:
            self.assertEqual(headers["type"], "Object"); self.assertNotIn("tdlmessage", payload)
            self.assertNotIn("collection", json.dumps(payload).casefold())

    @override_settings(GST_LOOKUP_PROVIDER="sandbox")
    def test_valid_gstin_without_sandbox_details_is_ready_with_a_gstin_fallback_warning(self):
        # Sandbox taxpayer lookup is optional enrichment, not an eligibility
        # gate: a valid GSTIN alone is a usable fallback party identity.
        result = party_eligibility(None, gstin="33AAACB2894G1ZJ")
        self.assertTrue(result["tally_ready"])
        self.assertEqual(result["party_name"], "33AAACB2894G1ZJ")
        self.assertEqual(result["name_source"], "Sandbox")
        self.assertEqual(result["warning_code"], "SANDBOX_PARTY_DETAILS_UNAVAILABLE")

    @override_settings(TALLY_ODBC_ENABLED=True, TALLY_ODBC_DSN="TestDSN", TALLY_ODBC_CONNECTION_STRING="")
    def test_odbc_company_and_exact_normalized_gstin_verification(self):
        class Cursor:
            def execute(self, query): self.query = query
            def fetchone(self): return ("ABC Enterprises", "Tamil Nadu", " 33aAacB2894g1zJ ")
        class Connection:
            def cursor(self): return Cursor()
            def close(self): pass
        status = odbc_company_status(lambda value, timeout: Connection())
        self.assertTrue(status["odbc_connected"])
        self.assertEqual(status["company_gstin"], "33AAACB2894G1ZJ")
        source_company = {"company_name": "ABC Enterprises"}
        self.assertTrue(verify_source_company("33AAACB2894G1ZJ", status, source_company)["company_verified"])
        mismatch = verify_source_company("32AAACB2894G1ZK", status, source_company)
        self.assertFalse(mismatch["company_verified"])
        self.assertEqual(mismatch["code"], "GSTIN_MISMATCH")

    @override_settings(TALLY_ODBC_ENABLED=True, TALLY_ODBC_DSN="TestDSN", TALLY_ODBC_CONNECTION_STRING="")
    def test_gstin_is_read_from_taxunit_when_company_field_is_blank(self):
        class Cursor:
            rows = []
            def execute(self, query):
                self.rows = [("GSTRCOMPANY", "Tamil Nadu", "")] if "FROM Company" in query else [("Tamil Nadu Registration", "33CSCPM4566P1Z7", "Tamil Nadu", "Regular")]
            def fetchall(self): return self.rows
        class Connection:
            def cursor(self): return Cursor()
            def close(self): pass
        status = odbc_company_status(lambda value, timeout: Connection(), requested_company=" gstrcompany ")
        self.assertEqual(status["company_gstin"], "33CSCPM4566P1Z7")
        self.assertEqual(status["gstin_read_source"], "ODBC_TAXUNIT")
        self.assertEqual(status["registration_type"], "Regular")

    @override_settings(TALLY_ODBC_ENABLED=False)
    def test_port_unreachable_is_reported_as_tally_not_reachable(self):
        with patch("gst_tally.tally.odbc.tcp_probe", return_value=(False, "TALLY_PORT_UNREACHABLE", "refused")):
            status = odbc_company_status()
        self.assertFalse(status["port_reachable"]); self.assertFalse(status["read_connected"])
        self.assertEqual(status["failure_type"], "TALLY_NOT_REACHABLE")
        self.assertEqual(status["company_name"], ""); self.assertEqual(status["company_gstin"], "")

    @override_settings(TALLY_WRITE_FORMAT="XML")
    def test_odbc_failure_does_not_hide_a_working_http_fallback(self):
        class Client:
            def post(self, payload, headers=None):
                if headers.get("subtype") == "Company":
                    return json.dumps({"status": "1", "tallymessage": [{"name": {"value": "SRI MAHALAKSHMI TRADERS"},
                        "statename": {"value": "Tamil Nadu"}, "gstregistrationdetails": [{"gstregnumber": {"value": "33AFHPM6103Q1Z8"}}]}]}).encode()
                return json.dumps({"status": "1", "tallymessage": [{}]}).encode()
        def failing_connect(value, timeout): raise RuntimeError("ODBC driver not found")
        with patch("gst_tally.tally.odbc.tcp_probe", return_value=(True, "", "")):
            status = odbc_company_status(failing_connect, requested_company="SRI MAHALAKSHMI TRADERS", read_client=Client())
        self.assertFalse(status["odbc_connected"])
        self.assertTrue(status["http_connected"])
        self.assertTrue(status["read_connected"], "ODBC failing must not hide a working HTTP read")
        self.assertEqual(status["company_gstin"], "33AFHPM6103Q1Z8")
        self.assertEqual(status["company_name"], "SRI MAHALAKSHMI TRADERS")
        self.assertEqual(status["failure_type"], "")
        self.assertIn("ODBC driver not found", status["odbc_error"])
        self.assertEqual(status["connection"], {
            "host": "127.0.0.1", "port": 9000, "port_reachable": True,
            "http_connected": True, "odbc_connected": False, "read_connected": True,
        })

    @override_settings(TALLY_WRITE_FORMAT="XML")
    def test_company_status_includes_active_tally_financial_period_from_sv_dates(self):
        class Client:
            def post(self, payload, headers=None):
                if headers and headers.get("subtype") == "Company":
                    return json.dumps({"status": "1", "tallymessage": [{"name": {"value": "SRI MAHALAKSHMI TRADERS"},
                        "statename": {"value": "Tamil Nadu"}, "gstregistrationdetails": [{"gstregnumber": {"value": "33AFHPM6103Q1Z8"}}]}]}).encode()
                if headers and headers.get("subtype") == "TaxUnit":
                    return json.dumps({"status": "1", "tallymessage": [{"gstregnumber": {"value": "33AFHPM6103Q1Z8"},
                        "statename": {"value": "Tamil Nadu"}, "gstregistrationdetails": [{"registrationtype": {"value": "Regular"}}]}]}).encode()
                assert b"SVFromDate" in payload
                assert b"SVToDate" in payload
                return b"""<ENVELOPE><HEADER><STATUS>1</STATUS></HEADER><BODY><DATA><COLLECTION>
                    <COMPANY><NAME>SRI MAHALAKSHMI TRADERS</NAME><STATENAME>Tamil Nadu</STATENAME>
                    <FINANCIALYEARFROM>1-Apr-25</FINANCIALYEARFROM><FINANCIALYEARTO>31-Mar-26</FINANCIALYEARTO>
                    </COMPANY></COLLECTION></DATA></BODY></ENVELOPE>"""
        def failing_connect(value, timeout): raise RuntimeError("ODBC driver not found")
        with patch("gst_tally.tally.odbc.tcp_probe", return_value=(True, "", "")):
            status = odbc_company_status(failing_connect, requested_company="SRI MAHALAKSHMI TRADERS", read_client=Client())

        self.assertEqual(status["financial_year_from"], "2025-04-01")
        self.assertEqual(status["financial_year_to"], "2026-03-31")
        self.assertEqual(status["financial_year"], "01 Apr 2025 - 31 Mar 2026")
        self.assertTrue(status["financial_year_available"])
        self.assertEqual(status["financial_year_error"], "")
        self.assertEqual(status["company_read"]["financial_year"], "01 Apr 2025 - 31 Mar 2026")

    @override_settings(TALLY_ODBC_ENABLED=True, TALLY_ODBC_DSN="TestDSN", TALLY_ODBC_CONNECTION_STRING="")
    def test_wrong_open_company_is_distinguished_from_no_company_open(self):
        class Cursor:
            def execute(self, query): self.query = query
            def fetchall(self): return [("A DIFFERENT COMPANY", "Kerala", "")]
        class Connection:
            def cursor(self): return Cursor()
            def close(self): pass
        status = odbc_company_status(lambda value, timeout: Connection(), requested_company="SRI MAHALAKSHMI TRADERS")
        self.assertFalse(status["company_detected"])
        self.assertEqual(status["failure_type"], "TALLY_COMPANY_NOT_OPEN")
        self.assertIn("A DIFFERENT COMPANY", status["open_companies"])

    def test_missing_tally_gstin_is_not_reported_as_mismatch(self):
        # The connection layer already knows this is specifically an unreadable
        # company-level GSTIN, not a generic "could not read anything" failure --
        # that specific code must surface instead of collapsing to GSTIN_NOT_READABLE.
        status = {"odbc_connected": True, "company_detected": True, "company_name": "Same Company", "company_gstin": "",
                  "failure_type": "TALLY_COMPANY_GSTIN_READ_FAILED", "message": "GSTIN unreadable"}
        result = verify_source_company("33CSCPM4566P1Z7", status, {"company_name": "Same Company"})
        self.assertEqual(result["verification"], "TALLY_COMPANY_GSTIN_READ_FAILED")
        self.assertIsNone(result["gstin_match"])

    def test_missing_tally_gstin_without_a_known_failure_type_falls_back_to_generic_code(self):
        status = {"odbc_connected": True, "company_detected": True, "company_name": "Same Company", "company_gstin": ""}
        result = verify_source_company("33CSCPM4566P1Z7", status, {"company_name": "Same Company"})
        self.assertEqual(result["verification"], "GSTIN_NOT_READABLE")
        self.assertIsNone(result["gstin_match"])

    def test_connection_failures_are_not_collapsed_into_gstin_not_readable(self):
        for failure_type in ("TALLY_NOT_REACHABLE", "TALLY_ODBC_UNAVAILABLE", "TALLY_HTTP_UNAVAILABLE",
                             "TALLY_COMPANY_NOT_DETECTED", "TALLY_COMPANY_NOT_OPEN"):
            with self.subTest(failure_type=failure_type):
                status = {"odbc_connected": False, "read_connected": False, "company_detected": False,
                          "company_name": "", "company_gstin": "", "failure_type": failure_type,
                          "message": f"diagnostic for {failure_type}"}
                result = verify_source_company("33CSCPM4566P1Z7", status, {"company_name": "Same Company"})
                self.assertFalse(result["company_verified"])
                self.assertEqual(result["verification"], failure_type)
                self.assertEqual(result["verification_code"], failure_type)
                self.assertIsNone(result["gstin_match"])
                self.assertEqual(result["message"], f"diagnostic for {failure_type}")

    def test_company_verification_prioritizes_exact_gstin_match_over_company_name(self):
        source = {"company_name": " sri  mahalakshmi traders, ", "pan": "WRONGPAN00", "state": "Kerala"}
        base = {"read_connected": True, "company_detected": True, "company_name": "SRI MAHALAKSHMI TRADERS.",
                "company_gstin": "33AFHPM6103Q1Z8", "pan": "DIFFERENT00", "state": "Tamil Nadu"}
        matched = verify_source_company("33AFHPM6103Q1Z8", base, source)
        self.assertTrue(matched["company_verified"]); self.assertEqual(matched["verification"], "MATCHED")
        self.assertTrue(matched["company_name_match"]); self.assertTrue(matched["gstin_match"])
        mismatch = verify_source_company("32AAACB2894G1ZK", base, source)
        self.assertFalse(mismatch["company_verified"]); self.assertEqual(mismatch["verification"], "GSTIN_MISMATCH")
        # A different display name must NOT override an exact GSTIN match.
        different_name = verify_source_company("33AFHPM6103Q1Z8", {**base, "company_name": "Different Company"}, source)
        self.assertTrue(different_name["company_verified"]); self.assertEqual(different_name["verification"], "MATCHED")
        self.assertFalse(different_name["company_name_match"])

    def test_acceptance_4_company_name_different_same_gstin_is_verified(self):
        # Fix-request Acceptance Test 4, verbatim values.
        source = {"company_name": "SRI MAHALAKSHMI TRADERS,"}
        status = {"read_connected": True, "company_detected": True, "company_name": "SRI MAHALAKSHMI TRADERS1",
                  "company_gstin": "33AFHPM6103Q1Z8"}
        result = verify_source_company("33AFHPM6103Q1Z8", status, source)
        self.assertTrue(result["company_verified"])
        self.assertEqual(result["verification_method"], "GSTIN")
        self.assertTrue(result["gstin_match"])
        self.assertFalse(result["company_name_match"])

    def test_acceptance_5_same_name_different_gstin_is_not_verified(self):
        # Fix-request Acceptance Test 5, verbatim values.
        source = {"company_name": "SRI MAHALAKSHMI TRADERS"}
        status = {"read_connected": True, "company_detected": True, "company_name": "SRI MAHALAKSHMI TRADERS",
                  "company_gstin": "33ABCDE1234F1Z5"}
        result = verify_source_company("33AFHPM6103Q1Z8", status, source)
        self.assertFalse(result["company_verified"])
        self.assertEqual(result["verification"], "GSTIN_MISMATCH")
        self.assertTrue(result["company_name_match"])

    def test_trailing_comma_in_source_company_name_does_not_cause_false_mismatch(self):
        # "SRI MAHALAKSHMI TRADERS," (as commonly copied from a source ledger) must
        # normalize the same as Tally's "SRI MAHALAKSHMI TRADERS" for name-match display.
        source = {"company_name": "SRI MAHALAKSHMI TRADERS,"}
        status = {"read_connected": True, "company_detected": True, "company_name": "SRI MAHALAKSHMI TRADERS",
                  "company_gstin": "33AFHPM6103Q1Z8"}
        result = verify_source_company("33AFHPM6103Q1Z8", status, source)
        self.assertTrue(result["company_verified"]); self.assertEqual(result["verification"], "MATCHED")
        self.assertTrue(result["company_name_match"])
        self.assertEqual(result["uploaded_company_name"], "SRI MAHALAKSHMI TRADERS,")
        self.assertEqual(result["tally_company_name"], "SRI MAHALAKSHMI TRADERS")

    def test_company_verification_is_identical_for_xml_and_json(self):
        source = {"company_name": "SRI MAHALAKSHMI TRADERS,"}
        same_gstin_different_name = {"read_connected": True, "company_detected": True,
                                     "company_name": "D", "company_gstin": "33AFHPM6103Q1Z8"}
        different_gstin = {**same_gstin_different_name, "company_gstin": "33CSCPM4566P1Z7"}
        same_gstin_same_name = {**same_gstin_different_name, "company_name": "SRI MAHALAKSHMI TRADERS"}
        for write_format in ("XML", "JSON"):
            with self.subTest(write_format=write_format), override_settings(TALLY_WRITE_FORMAT=write_format):
                result = verify_source_company("33AFHPM6103Q1Z8", same_gstin_different_name, source)
                self.assertTrue(result["company_verified"]); self.assertEqual(result["verification_code"], "MATCHED")
                mismatch = verify_source_company("33AFHPM6103Q1Z8", different_gstin, source)
                self.assertFalse(mismatch["company_verified"]); self.assertEqual(mismatch["verification_code"], "GSTIN_MISMATCH")
                matched = verify_source_company("33AFHPM6103Q1Z8", same_gstin_same_name, source)
                self.assertTrue(matched["company_verified"]); self.assertEqual(matched["verification_code"], "MATCHED")

    def test_active_company_verification_never_returns_old_identity_result(self):
        statuses = [
            {"read_connected": True, "company_detected": True, "company_name": "Same", "company_gstin": ""},
            {"read_connected": True, "company_detected": True, "company_name": "Same", "company_gstin": "33AFHPM6103Q1Z8"},
            {"read_connected": True, "company_detected": True, "company_name": "Other", "company_gstin": "33AFHPM6103Q1Z8"},
        ]
        for status in statuses:
            result = verify_source_company("33AFHPM6103Q1Z8", status, {"company_name": "Same"})
            self.assertNotIn(result["verification"], {"COMPANY_IDENTITY_MISMATCH", "MATCHED_BY_IDENTITY"})

    @override_settings(TALLY_ODBC_ENABLED=True, TALLY_ODBC_DSN="MissingDSN", TALLY_ODBC_CONNECTION_STRING="", TALLY_WRITE_FORMAT="XML")
    def test_company_gstin_read_uses_supported_http_taxunit_fallback_when_odbc_fails(self):
        class Client:
            def post(self, payload, headers=None):
                if headers["subtype"] == "Company":
                    return json.dumps({"status": "1", "tallymessage": [{
                        "name": {"value": "GSTRCOMPANY"}, "statename": {"value": "Tamil Nadu"},
                        "exciseunitname": {"value": "Tamil Nadu Registration"},
                        "gstregistrationnumber": {"value": ""},
                    }]}).encode()
                return json.dumps({"status": "1", "tallymessage": [{
                    "gstregnumber": {"value": "33CSCPM4566P1Z7"},
                    "statename": {"value": "Tamil Nadu"},
                }]}).encode()
        def failed_connect(value, timeout): raise RuntimeError("DSN missing")
        with patch("gst_tally.tally.odbc.tcp_probe", return_value=(True, "", "")):
            status = odbc_company_status(failed_connect, requested_company="GSTRCOMPANY",
                                         expected_gstin="33CSCPM4566P1Z7", read_client=Client())
        self.assertFalse(status["odbc_connected"])
        self.assertTrue(status["http_connected"])
        self.assertTrue(status["read_connected"])
        self.assertEqual(status["company_gstin"], "33CSCPM4566P1Z7")
        self.assertEqual(status["gstin_read_source"], "JSON_GST_REGISTRATION")
        result = verify_source_company("33CSCPM4566P1Z7", status, {"company_name": "GSTRCOMPANY"})
        self.assertTrue(result["company_verified"])
        self.assertEqual(result["verification"], "MATCHED")
        self.assertEqual(result["verification_diagnostics"], {
            "uploaded_company": "GSTRCOMPANY", "normalized_uploaded_company": "GSTRCOMPANY",
            "detected_company": "GSTRCOMPANY", "normalized_detected_company": "GSTRCOMPANY",
            "uploaded_gstin": "33CSCPM4566P1Z7", "tally_gstin": "33CSCPM4566P1Z7",
            "gstin_readable": True, "gstin_match": True, "company_name_match": True,
            "verification_code": "MATCHED",
        })

    def test_financial_year_start(self):
        self.assertEqual(get_financial_year_start(date(2026, 7, 15)), date(2026, 4, 1))
        self.assertEqual(get_financial_year_start(date(2027, 2, 15)), date(2026, 4, 1))
        self.assertEqual(get_financial_year_start(date(2027, 4, 2)), date(2027, 4, 1))
        self.assertEqual(financial_year_details(date(2026, 3, 31))["label"], "2025-26")
        self.assertEqual(financial_year_details(date(2026, 4, 1))["label"], "2026-27")

    def test_financial_year_boundary_dates(self):
        expected = {"31-03-2025": "2024-25", "01-04-2025": "2025-26",
                    "31-03-2026": "2025-26", "01-04-2026": "2026-27",
                    "31-03-2027": "2026-27", "01-04-2027": "2027-28"}
        for raw, label in expected.items():
            with self.subTest(raw=raw):
                self.assertEqual(financial_year_details(parse_date(raw))["label"], label)

    def test_xml_and_json_vouchers_set_period_from_invoice_date(self):
        value = voucher(18); value["invoice_date"] = "2025-04-01"; value["items"][0]["sales_ledger"] = "GST SALES @ 18%"
        xml = ET.fromstring(build_voucher(value, "D"))
        self.assertEqual(xml.findtext(".//SVFROMDATE"), "01-Apr-2025")
        self.assertEqual(xml.findtext(".//SVTODATE"), "31-Mar-2026")
        json_payload = build_json_voucher(value, "D")
        variables = {item["name"]: item["value"] for item in json_payload["static_variables"]}
        self.assertEqual(variables["svFromDate"], "20250401")
        self.assertEqual(variables["svToDate"], "20260331")

    def test_common_source_date_normalizer_is_day_first_and_supports_excel_serials(self):
        expected = date(2025, 4, 17)
        for value in ("17-04-2025", "17/04/2025", "17.04.2025", "2025-04-17", expected):
            self.assertEqual(parse_date(value), expected)
        self.assertEqual(parse_date(45764), expected)
        self.assertIsNone(parse_date("31-02-2025"))

    def test_same_invoice_date_normalizes_identically_for_excel_csv_and_json_and_xml(self):
        gstin = "33AAACB2894G1ZJ"
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["Customer GSTIN", "Invoice number", "Invoice Date", "Taxable Value", "Rate", "CGST", "SGST", "Invoice Value"])
        sheet.append([gstin, "INV-DATE", "17-04-2025", 1000, 18, 90, 90, 1180])
        excel = BytesIO(); excel.name = "dates.xlsx"; workbook.save(excel); excel.seek(0)
        excel_rows, _ = parse_gstr2b(excel)
        csv_file = BytesIO(b"GSTIN of supplier,Invoice number,Invoice Date,Taxable Value,Rate,Central Tax,State/UT Tax,Invoice Value\n33AAACB2894G1ZJ,INV-DATE,17-04-2025,1000,18,90,90,1180\n")
        csv_rows, _ = parse_gstr2a(csv_file)
        json_file = BytesIO(json.dumps({"b2b": [{"ctin": gstin, "inv": [{"inum": "INV-DATE", "idt": "17-04-2025", "val": 1180, "pos": "33", "itms": [{"itm_det": {"txval": 1000, "rt": 18, "camt": 90, "samt": 90}}]}]}]}).encode())
        json_rows, _ = parse_gstr1(json_file)
        self.assertEqual([excel_rows[0]["invoice_date"], csv_rows[0]["invoice_date"], json_rows[0]["invoice_date"]], [date(2025, 4, 17)] * 3)
        self.assertEqual(excel_rows[0]["source_line"]["source_row_number"], 2)
        self.assertEqual(excel_rows[0]["source_line"]["Invoice Value"], 1180)
        self.assertEqual(csv_rows[0]["source_line"]["source_row_number"], 2)
        self.assertEqual(csv_rows[0]["source_line"]["Invoice Value"], "1180")
        value = voucher(18); value["invoice_date"] = excel_rows[0]["invoice_date"].isoformat(); value["items"][0]["sales_ledger"] = "GST SALES @ 18%"
        xml = ET.fromstring(build_voucher(value, "GSTRCOMPANY"))
        self.assertEqual(xml.findtext(".//VOUCHER/DATE"), "20250417")

    def test_json_csv_and_excel_have_identical_financial_year_boundaries(self):
        gstin = "33AAACB2894G1ZJ"
        cases = (("31-03-2025", "2024-25"), ("01-04-2025", "2025-26"),
                 ("31-03-2026", "2025-26"), ("01-04-2026", "2026-27"),
                 ("31-03-2027", "2026-27"), ("01-04-2027", "2027-28"))
        for raw_date, expected_fy in cases:
            with self.subTest(raw_date=raw_date):
                workbook = Workbook(); sheet = workbook.active
                sheet.append(["Customer GSTIN", "Invoice number", "Invoice Date", "Taxable Value", "Rate", "CGST", "SGST", "Invoice Value"])
                sheet.append([gstin, "INV-FY", raw_date, 1000, 18, 90, 90, 1180])
                excel = BytesIO(); excel.name = "fy.xlsx"; workbook.save(excel); excel.seek(0)
                excel_rows, _ = parse_gstr2b(excel)
                csv_data = ("GSTIN of supplier,Invoice number,Invoice Date,Taxable Value,Rate,Central Tax,State/UT Tax,Invoice Value\n"
                            f"{gstin},INV-FY,{raw_date},1000,18,90,90,1180\n").encode()
                csv_rows, _ = parse_gstr2a(BytesIO(csv_data))
                json_data = {"b2b": [{"ctin": gstin, "inv": [{"inum": "INV-FY", "idt": raw_date, "val": 1180, "pos": "33",
                             "itms": [{"itm_det": {"txval": 1000, "rt": 18, "camt": 90, "samt": 90}}]}]}]}
                json_rows, _ = parse_gstr1(BytesIO(json.dumps(json_data).encode()))
                labels = [financial_year_details(rows[0]["invoice_date"])["label"] for rows in (excel_rows, csv_rows, json_rows)]
                self.assertEqual(labels, [expected_fy, expected_fy, expected_fy])

    def test_period_payload_detects_multiple_source_financial_years(self):
        result = _period_payload([{"invoice_date": "2026-03-31"}, {"invoice_date": "2026-04-01"}],
                                 {"books_from": date(2026, 4, 1), "ending_at": date(2027, 3, 31)})
        self.assertEqual(result["source_financial_years"], ["2025-26", "2026-27"])
        self.assertTrue(result["multiple_financial_years"]); self.assertEqual(result["tally_financial_year"], "2026-27")

    def test_company_address_and_state_normalization(self):
        company = normalize_company("33AAACB2894G1ZJ", {"TrdNm": "ABC Enterprises", "Addr1": "Madurai - 625011"}, {}, date(2026, 7, 15))
        self.assertEqual(company["state"], "Tamil Nadu"); self.assertEqual(company["country"], "India")
        self.assertEqual(company["pincode"], "625011"); self.assertNotIn("625011", company["address"])
        self.assertEqual(company["mobile"], ""); self.assertEqual(company["email"], "")

    def test_multiple_company_gstins_are_not_guessed(self):
        result = source_company_metadata(["33AAACB2894G1ZJ", "32AAACB2894G1ZK"])
        self.assertEqual(result["company_resolution_error"], "MULTIPLE_COMPANY_GSTINS")
        self.assertEqual(result["company_gstin"], "")

    def test_gstr1_top_level_gstin_is_company_and_tax_is_preserved(self):
        payload = {"gstin": "33AAACB2894G1ZJ", "fp": "072026", "b2b": [{"ctin": "32AAACB2894G1ZK", "inv": [{"inum": "INV-1", "idt": "15-07-2026", "val": 25020, "pos": "32", "itms": [{"itm_det": {"txval": 23828.70, "rt": 5, "iamt": 1191.44, "camt": 0, "samt": 0}}]}]}]}
        rows, metadata = parse_gstr1(BytesIO(json.dumps(payload).encode()))
        self.assertEqual(metadata["company_gstin"], "33AAACB2894G1ZJ")
        self.assertEqual(rows[0]["customer_gstin"], "32AAACB2894G1ZK")
        self.assertEqual(rows[0]["igst"], Decimal("1191.44")); self.assertEqual(rows[0]["sgst"], Decimal("0"))
    def test_tamil_nadu_party_is_intra_state(self): self.assertEqual(transaction_type("Tamil Nadu", "33")[0], "INTRA-STATE")
    def test_kerala_party_is_inter_state(self): self.assertEqual(transaction_type("Tamil Nadu", "32")[0], "INTER-STATE")
    def test_standard_intra_rates(self):
        for rate, half in ((5, "25.00"), (12, "60.00"), (18, "90.00"), (28, "140.00")):
            result = validate_voucher(voucher(rate)); self.assertTrue(result["valid"], result["errors"]); self.assertEqual(result["expected_cgst"], half)
    def test_interstate_igst(self):
        result = validate_voucher(voucher(18, "INTER-STATE")); self.assertTrue(result["valid"]); self.assertEqual(result["expected_igst"], "180.00")
    def test_mixed_rates(self):
        value = voucher(0); value["items"] = [{"taxable_value": "100", "gst_rate": "5"}, {"taxable_value": "200", "gst_rate": "12"}, {"taxable_value": "300", "gst_rate": "18"}]
        value.update(taxable_total="600", cgst="41.50", sgst="41.50", invoice_total="683")
        self.assertTrue(validate_voucher(value)["valid"])
    def test_roundoff_within_tolerance(self):
        value = voucher(18); value["invoice_total"] = "1180.01"; result = validate_voucher(value)
        self.assertTrue(result["valid"]); self.assertTrue(result["within_rounding_tolerance"])
        self.assertEqual(result["rounding_adjustment"], "0.01")

    def test_rounding_adjustment_balances_xml_and_json_payloads(self):
        value = voucher(18); value["invoice_total"] = "1180.01"; value["items"][0]["sales_ledger"] = "GST SALES @ 18%"
        value["rounding_adjustment"] = validate_voucher(value)["rounding_adjustment"]
        self.assertTrue(_voucher_balance(value)["balanced"])
        xml = ET.fromstring(build_voucher(value, "D"))
        xml_entries = {entry.findtext("LEDGERNAME"): entry.findtext("AMOUNT") for entry in xml.findall(".//LEDGERENTRIES.LIST")}
        self.assertEqual(xml_entries["Round Off"], "-0.01")
        json_payload = build_json_voucher(value, "D")
        json_entries = {entry["ledgername"]: entry["amount"] for entry in json_payload["tallymessage"][0]["ledgerentries"]}
        self.assertEqual(json_entries["Round Off"], "0.01")

    def test_purchase_negative_rounding_adjustment_serializes_with_balancing_positive_sign(self):
        value = voucher(18)
        value.update(voucher_type="Purchase", invoice_total="1179.99", rounding_adjustment="-0.01")
        value["rate_allocations"] = [{"gst_rate": "18", "sales_ledger": "GST Purchase 18%", "taxable_value": "1000"}]
        value["tax_allocations"] = [
            {"ledger": "Input CGST 9%", "amount": "90"},
            {"ledger": "Input SGST 9%", "amount": "90"},
        ]

        payload = build_json_voucher(value, "D")
        entries = payload["tallymessage"][0]["ledgerentries"]

        self.assertEqual(next(row["amount"] for row in entries if row["ledgername"] == "Round Off"), "0.01")
        self.assertEqual(sum(Decimal(row["amount"]) for row in entries), Decimal("0.00"))

    def test_unbalanced_voucher_is_detected_before_write(self):
        value = voucher(18); value["invoice_total"] = "1179.50"; value["items"][0]["sales_ledger"] = "GST SALES @ 18%"
        balance = _voucher_balance(value)
        self.assertFalse(balance["balanced"]); self.assertEqual(balance["difference"], "-0.50")
    def test_large_invoice_difference_is_not_rounding(self):
        value = voucher(18); value["invoice_total"] = "1108.00"; result = validate_voucher(value)
        self.assertFalse(result["within_rounding_tolerance"]); self.assertEqual(result["rounding_adjustment"], "0.00")

    def test_valid_single_rate_source_total(self):
        value = voucher(18); value.update(taxable_total="100000", cgst="9000", sgst="9000", invoice_total="118000")
        value["items"] = [{"taxable_value": "100000", "gst_rate": "18"}]
        result = validate_voucher(value)
        self.assertEqual(result["difference"], "0.00"); self.assertTrue(result["within_rounding_tolerance"])

    def test_valid_multi_rate_source_total(self):
        value = voucher(18); value.update(taxable_total="157267", cgst="16966.53", sgst="16966.53", invoice_total="191200.06")
        value["items"] = [{"taxable_value": "101017", "gst_rate": "18"}, {"taxable_value": "56250", "gst_rate": "28"}]
        result = validate_voucher(value)
        self.assertEqual(result["calculated_total"], "191200.06"); self.assertEqual(result["difference"], "0.00")

    def test_invalid_multi_rate_source_total_has_positive_variance(self):
        value = voucher(18); value.update(taxable_total="157267", cgst="16966.53", sgst="16966.53", invoice_total="119200.06")
        value["items"] = [{"taxable_value": "101017", "gst_rate": "18"}, {"taxable_value": "56250", "gst_rate": "28"}]
        result = validate_voucher(value)
        self.assertEqual(result["calculated_total"], "191200.06"); self.assertEqual(result["difference"], "72000.00")
        self.assertFalse(result["within_rounding_tolerance"])
    def test_missing_required_data(self):
        value = voucher(18); value["invoice_number"] = ""; self.assertFalse(validate_voucher(value)["valid"])
    def test_canonical_sales_ledger(self): self.assertEqual(sales_ledger_name("18.00"), "GST Sales 18%")
    def test_tally_http_body_rejection_is_not_success(self):
        result = parse_response("<ENVELOPE><BODY><DATA><IMPORTRESULT><CREATED>0</CREATED><ERRORS>1</ERRORS><LINEERROR>Invalid voucher</LINEERROR></IMPORTRESULT></DATA></BODY></ENVELOPE>")
        self.assertFalse(result.accepted); self.assertEqual(result.error, "Invalid voucher")
    def test_tally_exception_counter_is_an_explicit_rejection(self):
        result = parse_response("<RESPONSE><CREATED>0</CREATED><ERRORS>0</ERRORS><EXCEPTIONS>1</EXCEPTIONS></RESPONSE>")
        self.assertFalse(result.accepted); self.assertEqual(result.exceptions, 1)
        self.assertIn("Tally reported 1 exception(s)", result.error)

    def test_xml_exception_text_and_code_are_preserved(self):
        result = parse_response("<RESPONSE><EXCEPTIONS>1</EXCEPTIONS><EXCEPTIONDESC>Ledger GST SALES 18% does not exist</EXCEPTIONDESC><ERRORCODE>LEDGER_MISSING</ERRORCODE></RESPONSE>")
        self.assertEqual(result.error, "Ledger GST SALES 18% does not exist")
        self.assertEqual(result.exception_text, "Ledger GST SALES 18% does not exist")
        self.assertEqual(result.error_code, "LEDGER_MISSING")

    def test_live_style_out_of_range_line_error_is_preserved(self):
        raw = "<RESPONSE><LINEERROR>The date 1-4-2025 is Out of Range!</LINEERROR><CREATED>0</CREATED><ERRORS>0</ERRORS><EXCEPTIONS>1</EXCEPTIONS></RESPONSE>"
        result = parse_response(raw)
        self.assertFalse(result.accepted); self.assertEqual(result.created, 0)
        self.assertEqual(result.error, "The date 1-4-2025 is Out of Range!")

    def test_exceptions_without_any_known_diagnostic_tag_still_surfaces_nested_text(self):
        """Section 15: LINEERROR/EXCEPTIONDESC/DESCRIPTION are all empty (the
        exact 'EXCEPTIONS=1, LINEERROR unavailable' shape), but Tally put the
        real reason under a tag the structured search doesn't know by name --
        the last-resort scan must still surface it instead of discarding it."""
        raw = "<RESPONSE><CREATED>0</CREATED><ERRORS>0</ERRORS><EXCEPTIONS>1</EXCEPTIONS><REMARKS>Voucher date is beyond the last day of the active period</REMARKS></RESPONSE>"
        result = parse_response(raw)
        self.assertFalse(result.accepted)
        self.assertIn("REMARKS=Voucher date is beyond the last day of the active period", result.tally_exception_details)
        self.assertIn("Voucher date is beyond the last day of the active period", result.error)

    def test_exceptions_with_genuinely_no_nested_text_falls_back_to_the_generic_message(self):
        result = parse_response("<RESPONSE><CREATED>0</CREATED><ERRORS>0</ERRORS><EXCEPTIONS>1</EXCEPTIONS></RESPONSE>")
        self.assertEqual(result.tally_exception_details, "")
        self.assertEqual(result.error, "Tally reported 1 exception(s); full response captured for diagnosis")

    def test_json_nested_exception_text_and_code_are_preserved(self):
        raw = b'{"status":"0","data":{"import_result":{"exceptions":1,"exception":{"exceptionmessage":"Unsupported JSON voucher import","errorcode":"JSON_401"}}}}'
        result = parse_json_response(raw)
        self.assertEqual(result.error, "Unsupported JSON voucher import")
        self.assertEqual(result.exception_text, "Unsupported JSON voucher import")
        self.assertEqual(result.error_code, "JSON_401")

    def test_json_exception_without_text_reports_exact_counters_and_voucher_number(self):
        raw = b'{"status":"1","data":{"import_result":{"created":0,"altered":0,"ignored":0,"errors":0,"exceptions":1,"vchnumber":7}}}'

        result = parse_json_response(raw)

        self.assertFalse(result.accepted)
        self.assertEqual(result.error, "Tally rejected voucher 7: CREATED=0, ALTERED=0, IGNORED=0, ERRORS=0, EXCEPTIONS=1; no error text was returned by Tally.")

    def test_failed_voucher_diagnostics_include_request_and_all_tally_response_fields(self):
        value = voucher(18)
        value.update(invoice_number="7", voucher_type="Purchase", rate_allocations=[{
            "gst_rate": "18", "account_ledger": "GST Purchase 18%", "taxable_value": "1000"}])
        response = TallyResponse(created=0, altered=0, ignored=0, errors=0, exceptions=1,
                                 error="Tally rejected voucher 7", raw='{"exceptions":1}')
        request_payload = {"tallymessage": [{"vouchernumber": "7"}]}

        result = _voucher_diagnostics(value, {"company_name": "D", "company_gstin": "33AFHPM6103Q1Z8"},
                                      {}, response, request_payload=request_payload)

        self.assertEqual(result["request_payload"], request_payload)
        self.assertEqual(len(result["request_payload_hash"]), 64)
        self.assertEqual(result["voucher_data_summary"]["invoice_no"], "7")
        self.assertEqual(result["voucher_data_summary"]["voucher_type"], "Purchase")
        self.assertEqual(result["voucher_data_summary"]["purchase_ledgers"], ["GST Purchase 18%"])
        self.assertEqual(result["tally_response_fields"], {
            "created": 0, "altered": 0, "ignored": 0, "errors": 0, "exceptions": 1,
            "line_error": "", "description": "", "exception_text": "", "error_code": ""})
        self.assertEqual(_write_error_code(response), "TALLY_EXCEPTION_WITHOUT_MESSAGE")

    @override_settings(TALLY_DRY_RUN=False, TALLY_EXPECTED_COMPANY="")
    def test_company_parser_ignores_cmpinfo_counter(self):
        class Client:
            def post(self, payload):
                return b'<ENVELOPE><HEADER><STATUS>1</STATUS></HEADER><BODY><DESC><CMPINFO><COMPANY>0</COMPANY></CMPINFO></DESC><DATA><COLLECTION><COMPANY NAME="GSTRCOMPANY"><NAME>GSTRCOMPANY</NAME><STATENAME>Tamil Nadu</STATENAME><GSTREGISTRATIONNUMBER>33ABCDE1234F1Z5</GSTREGISTRATIONNUMBER></COMPANY></COLLECTION></DATA></BODY></ENVELOPE>'
        result = connection_status(Client())
        self.assertTrue(result["reachable"]); self.assertTrue(result["company_open"])
        self.assertEqual(result["company"], "GSTRCOMPANY"); self.assertEqual(result["state"], "Tamil Nadu")

    @override_settings(TALLY_DRY_RUN=False)
    def test_read_timeout_is_classified(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            def post(self, payload): raise TallyConnectionError("TALLY_READ_TIMEOUT", "No HTTP response")
        result = connection_status(Client())
        self.assertTrue(result["tcp_connected"]); self.assertFalse(result["tally_response_received"])
        self.assertEqual(result["failure_type"], "TALLY_READ_TIMEOUT")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=True, TALLY_VERSION="7.0")
    def test_step3_reports_http_connected_when_odbc_fails(self):
        # TALLY_VERSION="7.0": this stub only speaks XML (no JSON handling),
        # simulating a pre-7.1 Tally so the version-aware router deterministically
        # selects the XML company read this test exercises.
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><BODY><DATA><COLLECTION><COMPANY NAME="GSTRCOMPANY"><NAME>GSTRCOMPANY</NAME><STATENAME>Tamil Nadu</STATENAME><GSTREGISTRATIONNUMBER>33ABCDE1234F1Z5</GSTREGISTRATIONNUMBER></COMPANY></COLLECTION></DATA></BODY></ENVELOPE>'

        def failed_connect(value, timeout):
            raise RuntimeError("ODBC driver not found")

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client(), odbc_connect=failed_connect)

        self.assertTrue(result["read_connected"])
        self.assertTrue(result["http_connected"])
        self.assertFalse(result["odbc_connected"])
        self.assertTrue(result["company_open"])
        self.assertEqual(result["host"], "127.0.0.1")
        self.assertEqual(result["port"], 9000)
        self.assertEqual(result["error_code"], "TALLY_ODBC_UNAVAILABLE")
        self.assertIn("ODBC driver not found", result["error_message"])

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=True)
    def test_step3_tcp_refused_does_not_attempt_http_or_odbc(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            def post(self, payload):
                raise AssertionError("HTTP must not be attempted when raw TCP fails")

        def odbc_connect(value, timeout):
            raise AssertionError("ODBC must not be attempted when raw TCP fails")

        with patch("gst_tally.tally.connection.tcp_probe",
                   return_value=(False, "TALLY_CONNECTION_REFUSED", "Tally refused the TCP connection")), \
             patch("gst_tally.tally.connection.local_port_listening", return_value=None):
            result = step3_connection_check(Client(), odbc_connect=odbc_connect)

        self.assertFalse(result["tcp_connected"])
        self.assertFalse(result["http_connected"])
        self.assertFalse(result["odbc_connected"])
        self.assertFalse(result["read_connected"])
        self.assertFalse(result["company_open"])
        self.assertFalse(result["can_import"])
        self.assertEqual(result["error_code"], "TALLY_CONNECTION_REFUSED")
        self.assertEqual(result["http_error"], "Tally refused the TCP connection")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=True)
    def test_step3_reports_local_port_not_listening_when_netstat_has_no_listener(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            def post(self, payload):
                raise AssertionError("HTTP must not be attempted when local port is closed")

        with patch("gst_tally.tally.connection.tcp_probe",
                   return_value=(False, "TALLY_CONNECTION_REFUSED", "Tally refused the TCP connection")), \
             patch("gst_tally.tally.connection.local_port_listening", return_value=False):
            result = step3_connection_check(Client())

        self.assertFalse(result["tcp_connected"])
        self.assertEqual(result["error_code"], "TALLY_PORT_NOT_LISTENING")
        self.assertEqual(result["error_message"], "No process is listening on 127.0.0.1:9000.")
        self.assertFalse(result["can_import"])

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=True, TALLY_WRITE_FORMAT="XML", TALLY_VERSION="7.0")
    def test_step3_reports_missing_odbc_driver_without_treating_http_as_import_ready(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><BODY><DATA><COLLECTION><COMPANY NAME="GSTRCOMPANY"><NAME>GSTRCOMPANY</NAME><STATENAME>Tamil Nadu</STATENAME><GSTREGISTRATIONNUMBER>33ABCDE1234F1Z5</GSTREGISTRATIONNUMBER></COMPANY></COLLECTION></DATA></BODY></ENVELOPE>'

        def missing_driver(value, timeout):
            raise RuntimeError("IM002 - Data source name not found and no default driver specified")

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client(), odbc_connect=missing_driver)

        self.assertTrue(result["tcp_connected"])
        self.assertTrue(result["http_connected"])
        self.assertFalse(result["odbc_connected"])
        self.assertTrue(result["read_connected"])
        self.assertTrue(result["company_open"])
        self.assertFalse(result["can_import"])
        self.assertEqual(result["error_code"], "ODBC_DRIVER_NOT_AVAILABLE")
        self.assertIn("IM002", result["odbc_error"])

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=True, TALLY_WRITE_FORMAT="JSON", TALLY_VERSION="7.0")
    def test_step3_can_continue_when_http_company_is_open_and_odbc_driver_is_missing(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><BODY><DATA><COLLECTION><COMPANY NAME="GSTRCOMPANY"><NAME>GSTRCOMPANY</NAME><STATENAME>Tamil Nadu</STATENAME></COMPANY></COLLECTION></DATA></BODY></ENVELOPE>'

        def missing_driver(value, timeout):
            raise RuntimeError("IM002 - Data source name not found and no default driver specified")

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client(), odbc_connect=missing_driver)

        self.assertTrue(result["tcp_connected"])
        self.assertTrue(result["http_connected"])
        self.assertTrue(result["company_open"])
        self.assertTrue(result["can_import"])
        self.assertFalse(result["odbc_connected"])
        self.assertEqual(result["error_code"], "ODBC_DRIVER_NOT_AVAILABLE")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=True, TALLY_WRITE_FORMAT="JSON")
    def test_step3_uses_json_transport_and_ignores_odbc_when_tally_is_7_1_or_above(self):
        """PART 13 Test F: Tally >= 7.1, ODBC false, JSON company read succeeds
        -> company_open must be true and transport must honestly be JSON."""
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return json.dumps({"status": "1", "prod_maj_rel": "7", "prod_min_rel": "1", "tallymessage": [{
                    "name": "GSTRCOMPANY", "statename": "Tamil Nadu",
                    "gstregistrationnumber": "33ABCDE1234F1Z5"}]}).encode()

        def missing_driver(value, timeout):
            raise RuntimeError("IM002 - Data source name not found and no default driver specified")

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client(), odbc_connect=missing_driver)

        self.assertTrue(result["tcp_connected"])
        self.assertTrue(result["http_connected"])
        self.assertTrue(result["company_open"])
        self.assertTrue(result["can_import"])
        self.assertFalse(result["odbc_connected"])
        self.assertEqual(result["tally_version_detected"], "7.1")
        self.assertEqual(result["version_source"], "DETECTED")
        self.assertEqual(result["requested_transport"], "JSON")
        self.assertEqual(result["actual_transport"], "JSON")
        self.assertEqual(result["company_name"], "GSTRCOMPANY")
        self.assertEqual(result["company_gstin"], "33ABCDE1234F1Z5")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False, TALLY_VERSION="7.0")
    def test_step3_reads_company_gstin_from_json_registration_when_company_collection_omits_it(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><BODY><DATA><COLLECTION><COMPANY NAME="GSTRCOMPANY"><NAME>GSTRCOMPANY</NAME><STATENAME>Tamil Nadu</STATENAME></COMPANY></COLLECTION></DATA></BODY></ENVELOPE>'

        registration = {"gstin": "33ABCDE1234F1Z5", "state": "Tamil Nadu",
                        "company_name": "GSTRCOMPANY", "trade_name": "GSTRCOMPANY"}
        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")), \
             patch("gst_tally.tally.connection.get_tally_company_gst_registration", return_value=registration):
            result = step3_connection_check(Client())

        self.assertTrue(result["http_connected"])
        self.assertTrue(result["company_open"])
        self.assertEqual(result["company_gstin"], "33ABCDE1234F1Z5")
        self.assertEqual(result["gstin"], "33ABCDE1234F1Z5")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False, TALLY_EXPECTED_COMPANY="", TALLY_VERSION="7.0")
    def test_step3_reports_company_not_open_when_http_server_has_no_company(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><HEADER><STATUS>1</STATUS></HEADER><BODY><DATA><COLLECTION></COLLECTION></DATA></BODY></ENVELOPE>'

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client())

        self.assertTrue(result["read_connected"])
        self.assertFalse(result["company_open"])
        self.assertEqual(result["error_code"], "COMPANY_NOT_OPEN")
        self.assertEqual(result["error_message"], "Tally is connected, but no company is open.")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False)
    def test_step3_normalizes_refused_and_timeout_codes(self):
        class Client:
            base_url = "http://127.0.0.1:9000"
            def __init__(self, code): self.code = code
            def post(self, payload, headers=None): raise TallyConnectionError(self.code, "raw failure")

        with patch("gst_tally.tally.connection.tcp_probe",
                   return_value=(False, "TALLY_CONNECTION_REFUSED", "Tally refused the TCP connection")), \
             patch("gst_tally.tally.connection.local_port_listening", return_value=None):
            refused = step3_connection_check(Client("TALLY_CONNECTION_REFUSED"))
        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            timed_out = step3_connection_check(Client("TALLY_READ_TIMEOUT"))

        self.assertEqual(refused["error_code"], "TALLY_CONNECTION_REFUSED")
        self.assertFalse(refused["read_connected"])
        self.assertEqual(timed_out["error_code"], "TALLY_CONNECTION_TIMEOUT")
        self.assertFalse(timed_out["read_connected"])

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False, TALLY_VERSION="")
    def test_step3_falls_back_to_xml_when_version_is_unknown_and_xml_works(self):
        """Version fully unknown (no live probe, no usable configured
        fallback) must never hard-fail the connection: JSON isn't preferred
        without a supporting version, so this goes straight to XML, and a
        working XML read is a fully successful connection."""
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><BODY><DATA><COLLECTION><COMPANY NAME="GSTRCOMPANY"><NAME>GSTRCOMPANY</NAME><STATENAME>Tamil Nadu</STATENAME><GSTREGISTRATIONNUMBER>33ABCDE1234F1Z5</GSTREGISTRATIONNUMBER></COMPANY></COLLECTION></DATA></BODY></ENVELOPE>'

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client())

        self.assertEqual(result["error_code"], "")
        self.assertFalse(result["json_supported_by_version"])
        self.assertEqual(result["requested_transport"], "XML")
        self.assertEqual(result["actual_transport"], "XML")
        self.assertFalse(result["fallback_used"])
        self.assertTrue(result["read_connected"])
        self.assertTrue(result["company_open"])

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False, TALLY_VERSION="7.1")
    def test_step3_prefers_json_but_falls_back_to_xml_when_json_read_fails(self):
        """PART 25 TEST 2 / TEST 4 -- and the exact bug this turn fixes:
        configured version says JSON is eligible, but this stub only speaks
        XML (the live JSON probe fails to parse) -- must NOT fail the whole
        connection. XML is tried automatically and, since it succeeds, the
        connection is fully readable with an honestly reported fallback."""
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><BODY><DATA><COLLECTION><COMPANY NAME="GSTRCOMPANY"><NAME>GSTRCOMPANY</NAME><STATENAME>Tamil Nadu</STATENAME><GSTREGISTRATIONNUMBER>33ABCDE1234F1Z5</GSTREGISTRATIONNUMBER></COMPANY></COLLECTION></DATA></BODY></ENVELOPE>'

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client())

        self.assertEqual(result["error_code"], "")
        self.assertTrue(result["json_supported_by_version"])
        self.assertFalse(result["json_capability_confirmed"])
        self.assertFalse(result["json_connected"])
        self.assertTrue(result["xml_connected"])
        self.assertEqual(result["requested_transport"], "JSON")
        self.assertEqual(result["actual_transport"], "XML")
        self.assertTrue(result["fallback_used"])
        self.assertIn("XML HTTP integration used", result["fallback_reason"])
        self.assertTrue(result["read_connected"])
        self.assertTrue(result["http_connected"])
        self.assertTrue(result["company_open"])
        self.assertEqual(result["company_name"], "GSTRCOMPANY")
        self.assertEqual(result["company_gstin"], "33ABCDE1234F1Z5")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False, TALLY_VERSION="7.1")
    def test_step3_json_fails_xml_open_company_field_is_empty(self):
        """PART 25 TEST 5: JSON fails, XML succeeds but reports no open
        company -> COMPANY_NOT_OPEN, and specifically NOT a transport error --
        transport worked fine, it's the company that isn't open."""
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def post(self, payload, headers=None):
                return b'<ENVELOPE><HEADER><STATUS>1</STATUS></HEADER><BODY><DATA><COLLECTION></COLLECTION></DATA></BODY></ENVELOPE>'

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client())

        self.assertTrue(result["xml_connected"])
        self.assertTrue(result["read_connected"])
        self.assertFalse(result["company_open"])
        self.assertEqual(result["error_code"], "COMPANY_NOT_OPEN")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False, TALLY_VERSION="7.1")
    def test_step3_json_and_xml_both_fail_reports_transport_unavailable(self):
        """PART 25 TEST 3: neither transport can read Tally -> read_connected
        is false and a specific TALLY_TRANSPORT_UNAVAILABLE-style code is
        returned (never a bare COMPANY_NOT_OPEN, which would misleadingly
        imply transport itself was fine)."""
        class Client:
            base_url = "http://127.0.0.1:9000"
            def post(self, payload, headers=None):
                raise TallyConnectionError("TALLY_INVALID_RESPONSE", "neither transport responded usefully")

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            result = step3_connection_check(Client())

        self.assertFalse(result["read_connected"])
        self.assertFalse(result["json_connected"])
        self.assertFalse(result["xml_connected"])
        self.assertFalse(result["company_open"])
        self.assertEqual(result["error_code"], "TALLY_TRANSPORT_UNAVAILABLE")
        self.assertNotEqual(result["error_code"], "COMPANY_NOT_OPEN")

    @override_settings(TALLY_DRY_RUN=False, TALLY_ODBC_ENABLED=False, TALLY_WRITE_FORMAT="JSON")
    def test_step3_retry_reflects_a_company_switch_with_no_stale_cache(self):
        """PART 11/12/28: a fresh call must reflect whatever Tally reports
        right now -- switching companies between two calls (as a user
        clicking Retry after switching in Tally would) must never replay a
        cached result from the first call."""
        class Client:
            base_url = "http://127.0.0.1:9000"
            last_http_status = 200
            def __init__(self, company_name):
                self.company_name = company_name
            def post(self, payload, headers=None):
                return json.dumps({"status": "1", "prod_maj_rel": "7", "prod_min_rel": "1", "tallymessage": [{
                    "name": self.company_name, "statename": "Tamil Nadu",
                    "gstregistrationnumber": "33ABCDE1234F1Z5"}]}).encode()

        with patch("gst_tally.tally.connection.tcp_probe", return_value=(True, "", "")):
            first = step3_connection_check(Client("Company A"))
            second = step3_connection_check(Client("Company B"))

        self.assertEqual(first["company_name"], "Company A")
        self.assertEqual(second["company_name"], "Company B")
