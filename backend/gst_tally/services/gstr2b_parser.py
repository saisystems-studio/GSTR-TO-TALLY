from openpyxl import Workbook, load_workbook
from gst_tally.utils.file_utils import normalize_header, first_value, parse_date, parse_decimal, json_safe
from gst_tally.utils.uploaded_files import validate_upload_type
from .source_party import add_source_party, log_source_parties

ALIASES = {
 "gstin": ["Customer GSTIN", "Customer GSTIN/UIN", "GSTIN/UIN of Customer", "GSTIN of supplier", "GSTIN/UIN of Supplier", "Supplier GSTIN", "ctin", "GSTIN"], "party_name": ["Trade/Legal name", "Supplier Name", "Customer Name", "Party Name"],
 "invoice_number": ["Invoice number", "Invoice No", "inum"], "invoice_date": ["Invoice Date", "idt"],
 "taxable_value": ["Taxable Value", "txval"], "cgst": ["Central Tax", "CGST"], "sgst": ["State/UT Tax", "SGST"],
 "igst": ["Integrated Tax", "IGST"], "cess": ["Cess"], "total_value": ["Invoice Value", "Total Value"],
 "itc_value": ["ITC Available", "ITC Value", "Total ITC Available"],
 "tax_percent": ["Rate", "Tax %", "rt"], "state_code": ["Place Of Supply", "State Code", "pos"],
 "reverse_charge": ["Reverse Charge", "rchrg"], "invoice_type": ["Invoice Type", "inv_typ"],
 "filing_period": ["Filing Period", "Return Period"], "filing_date": ["Filing Date"],
}
COMPANY_ALIASES = ["GSTIN of recipient", "Recipient GSTIN", "Company GSTIN", "My GSTIN", "Taxpayer GSTIN"]
def _load_source_workbook(file_obj):
    extension = validate_upload_type(file_obj, {"xlsx", "xls"})
    if extension != "xls":
        return load_workbook(file_obj, read_only=True, data_only=True)
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("Legacy .xls support requires the xlrd package") from exc
    raw = file_obj.read()
    legacy = xlrd.open_workbook(file_contents=raw)
    workbook = Workbook(); workbook.remove(workbook.active)
    for legacy_sheet in legacy.sheets():
        sheet = workbook.create_sheet(legacy_sheet.name)
        for row_index in range(legacy_sheet.nrows):
            values = []
            for cell in legacy_sheet.row(row_index):
                value = xlrd.xldate_as_datetime(cell.value, legacy.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                values.append(value)
            sheet.append(values)
    return workbook
def _sheet_and_header(workbook):
    required = {normalize_header(x) for x in ALIASES["invoice_number"]}
    for sheet in workbook.worksheets:
        for index, values in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30), values_only=True), 1):
            normalized = {normalize_header(v) for v in values if v is not None}
            if normalized & required:
                return sheet, index, list(values)
    raise ValueError("No supported invoice sheet/header was found")
def parse(file_obj):
    try:
        workbook = _load_source_workbook(file_obj)
        sheet, header_row, headers = _sheet_and_header(workbook)
        rows, parties, company_candidates, company_source = [], {}, [], {}
        normalized_headers = {normalize_header(value) for value in headers}
        company_aliases = [*COMPANY_ALIASES, "GSTIN"] if normalize_header("Customer GSTIN") in normalized_headers else COMPANY_ALIASES
        for source_row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            source = {str(headers[i] or f"column_{i}"): values[i] for i in range(min(len(headers), len(values)))}
            if not any(str(v).strip() for v in source.values() if v is not None): continue
            row = {k: first_value(source, v) for k, v in ALIASES.items()}
            company_value = first_value(source, company_aliases)
            if company_value: company_candidates.append(company_value); company_source = source
            add_source_party(parties, row["gstin"], source)
            rows.append({"customer_gstin": str(row["gstin"] or "").strip().upper(),
                "invoice_no": str(row["invoice_number"] or "").strip(), "invoice_date": parse_date(row["invoice_date"]),
                "taxable_value": parse_decimal(row["taxable_value"]), "tax_percent": parse_decimal(row["tax_percent"]),
                "cgst": parse_decimal(row["cgst"]), "sgst": parse_decimal(row["sgst"]), "igst": parse_decimal(row["igst"]),
                "invoice_value": parse_decimal(row["total_value"]), "state_code": str(row["state_code"] or "").strip(),
                "reverse_charge": str(row["reverse_charge"] or "").strip(), "invoice_type": str(row["invoice_type"] or "").strip(),
                "filing_period": str(row["filing_period"] or "").strip(), "filing_type": "GSTR2B",
                "filing_date": parse_date(row["filing_date"]), "source_type": str(sheet.title or "B2B").upper(),
                "source_line": json_safe({"source_row_number": source_row_number, **source})})
        log_source_parties(parties)
        from .company import source_company_metadata
        return rows, {"parties": parties, **source_company_metadata(company_candidates, company_source)}
    except ValueError: raise
    except Exception as exc: raise ValueError(f"Invalid Excel file: {exc}") from exc
