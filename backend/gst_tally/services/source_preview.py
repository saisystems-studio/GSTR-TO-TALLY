import json
import logging
from collections import OrderedDict
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from gst_tally.utils.uploaded_files import file_type_label, load_json_upload, validate_upload_type
from .csv_preview_parser import parse_csv_preview

RETURN_TYPES = {"GSTR1", "GSTR2A", "GSTR2B"}
ALLOWED_EXTENSIONS = {"json", "csv", "xlsx", "xls"}
logger = logging.getLogger(__name__)


def _display(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def _flatten(value, prefix="", output=None):
    output = output if output is not None else OrderedDict()
    if isinstance(value, dict):
        for key, child in value.items():
            name = f"{prefix}.{key}" if prefix else str(key)
            if isinstance(child, dict):
                _flatten(child, name, output)
            elif isinstance(child, list):
                output[name] = json.dumps(child, ensure_ascii=False, separators=(",", ":"))
            else:
                output[name] = _display(child)
    else:
        output[prefix or "value"] = _display(value)
    return output


def _ci_get(mapping, *aliases, default=None):
    if not isinstance(mapping, dict):
        return default
    keys = {str(key).casefold(): key for key in mapping}
    for alias in aliases:
        actual = keys.get(alias.casefold())
        if actual is not None:
            return mapping[actual]
    return default


def _flatten_except(mapping, excluded):
    output = OrderedDict()
    excluded = {key.casefold() for key in excluded}
    for key, value in mapping.items():
        if str(key).casefold() not in excluded:
            _flatten(value, str(key), output)
    return output


def _json_rows(payload):
    # GSTR-1 B2B has party -> invoice -> item nesting. Expand only arrays into rows;
    # retain original key names and values without GST calculations or normalization.
    b2b = _ci_get(payload, "b2b", default=None) if isinstance(payload, dict) else None
    if isinstance(payload, dict) and isinstance(b2b, list):
        rows = []
        root_values = _flatten_except(payload, {"b2b"})
        invoice_count = 0
        item_count = 0
        for party in b2b:
            party_values = _flatten_except(party, {"inv"}) if isinstance(party, dict) else OrderedDict()
            invoices = _ci_get(party, "inv", "invoices", default=[])
            if not invoices:
                rows.append(OrderedDict((*root_values.items(), *party_values.items())))
                continue
            for invoice in invoices:
                invoice_count += 1
                invoice_values = _flatten_except(invoice, {"itms", "items"}) if isinstance(invoice, dict) else OrderedDict()
                items = _ci_get(invoice, "itms", "items", default=[])
                if not items:
                    rows.append(OrderedDict((*root_values.items(), *party_values.items(), *invoice_values.items())))
                    continue
                for item in items:
                    item_count += 1
                    row = OrderedDict((*root_values.items(), *party_values.items(), *invoice_values.items()))
                    row.update(_flatten(item))
                    rows.append(row)
        logger.info("GST JSON preview counts - invoices: %d, item entries: %d, preview rows: %d",
                    invoice_count, item_count, len(rows))
        return rows
    if isinstance(payload, list):
        return [_flatten(item) for item in payload]
    if isinstance(payload, dict):
        for value in payload.values():
            if isinstance(value, list) and value:
                return [_flatten(item) for item in value]
        return [_flatten(payload)]
    return [_flatten(payload)]


def _matrix(rows):
    columns = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    return columns, [[row.get(column, "") for column in columns] for row in rows]


def _structured_headers(columns):
    groups = []
    invoice = {"inum", "inv_typ", "idt", "val"}
    tax = {"itm_det.iamt", "itm_det.camt", "itm_det.samt", "itm_det.csamt"}
    item = {"num", "itm_det.rt", "itm_det.txval"}
    labels = {"invoice": "Invoice Details", "item": "Item Details", "tax": "Tax Amount"}
    kinds = []
    for column in columns:
        kinds.append("invoice" if column in invoice else "tax" if column in tax else "item" if column in item else None)
    index = 0
    while index < len(columns):
        kind = kinds[index]
        if kind is None:
            groups.append({"label": columns[index], "row_span": 2})
            index += 1
            continue
        children = []
        while index < len(columns) and kinds[index] == kind:
            children.append({"label": columns[index]})
            index += 1
        groups.append({"label": labels[kind], "children": children})
    return groups


def _preview_json(file_obj):
    payload = load_json_upload(file_obj, object_pairs_hook=OrderedDict, parse_int=str, parse_float=str)
    columns, rows = _matrix(_json_rows(payload))
    return {"sheet_name": None, "sheets": [], "columns": columns, "rows": rows,
            "layout_mode": "structured", "title": "Goods and Services Tax - GSTR-1 (B2B)",
            "section_title": "B2B Invoice Source Data", "headers": _structured_headers(columns),
            "merged_cells": [], "column_widths": []}


def _preview_excel(file_obj, requested_sheet=None):
    try:
        if Path(getattr(file_obj, "name", "")).suffix.lower() == ".xls":
            from .gstr2b_parser import _load_source_workbook
            workbook = _load_source_workbook(file_obj)
        else:
            workbook = load_workbook(file_obj, read_only=False, data_only=True)
    except Exception as exc:
        raise ValueError(f"Invalid Excel file: {exc}") from exc
    sheets = workbook.sheetnames
    if not sheets:
        return {"sheet_name": None, "sheets": [], "columns": [], "rows": [], "layout_mode": "worksheet", "merged_cells": [], "column_widths": []}
    sheet_name = requested_sheet if requested_sheet in sheets else sheets[0]
    sheet = workbook[sheet_name]
    values = [[_display(cell.value) for cell in row] for row in sheet.iter_rows()]
    while values and not any(value != "" for value in values[-1]):
        values.pop()
    if not values:
        return {"sheet_name": sheet_name, "sheets": sheets, "columns": [], "rows": [], "layout_mode": "worksheet", "merged_cells": [], "column_widths": []}
    width = max(len(row) for row in values)
    rows = [list(row) + [""] * (width - len(row)) for row in values]
    merged = [{"start_row": item.min_row - 1, "start_col": item.min_col - 1, "row_span": item.max_row - item.min_row + 1, "col_span": item.max_col - item.min_col + 1} for item in sheet.merged_cells.ranges]
    widths = [sheet.column_dimensions[get_column_letter(index + 1)].width or 13 for index in range(width)]
    detected_header = next((index for index, row in enumerate(rows)
                            if any(value.strip().casefold() == "gstin of supplier" for value in row)), None)
    data_start = detected_header + 1 if detected_header is not None else _excel_header_row_count(rows)
    candidates = rows[:data_start]
    header_index = detected_header if detected_header is not None else (max(range(len(candidates)), key=lambda index: sum(bool(value.strip()) for value in candidates[index])) if candidates else None)
    columns = candidates[header_index] if header_index is not None else (rows[0] if rows else [])
    title_rows = [row for index, row in enumerate(candidates) if index != header_index and sum(bool(value.strip()) for value in row) == 1]
    titles = [next(value for value in row if value.strip()) for row in title_rows]
    data_rows = rows[data_start:]
    return {"sheet_name": sheet_name, "sheets": sheets, "columns": columns, "rows": data_rows,
            "layout_mode": "worksheet", "title": titles[0] if titles else "",
            "section_title": titles[1] if len(titles) > 1 else "", "headers": [],
            "merged_cells": [], "column_widths": widths, "header_row_count": 0}


def _excel_header_row_count(rows):
    for index, row in enumerate(rows):
        values = [value.strip() for value in row if value.strip()]
        if not values:
            continue
        numeric = sum(value.replace(",", "").replace(".", "", 1).lstrip("+-").isdigit() for value in values)
        identifiers = sum(any(char.isdigit() for char in value) and len(value) >= 10 for value in values)
        if index > 0 and (numeric >= max(2, len(values) // 3) or identifiers >= 2):
            return index
    return 1 if rows else 0


def preview_file(file_obj, return_type, sheet_name=None):
    if return_type not in RETURN_TYPES:
        raise ValueError("Unsupported return type")
    extension = validate_upload_type(file_obj, ALLOWED_EXTENSIONS)
    if extension == "json":
        result = _preview_json(file_obj)
    elif extension == "csv":
        result = parse_csv_preview(file_obj)
        result.update({"layout_mode": "worksheet", "headers": [], "merged_cells": [], "column_widths": []})
    else:
        result = _preview_excel(file_obj, sheet_name)
    return {"return_type": return_type, "file_name": file_obj.name, "file_type": file_type_label(extension),
            **result, "row_count": len(result["rows"])}
